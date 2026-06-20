"""
Coach Daily Notes API Routes

Lets supervisors (with the ``coach-notes`` permission) and admins write daily
evaluation notes about coaches — free text plus three preset rated categories
(punctuality / member handling / performance). Notes accumulate and can be
reviewed/exported as a monthly report per coach. These notes are private to
supervisors+admins (never exposed to the coach or member portal).
"""
from fastapi import APIRouter, HTTPException, Depends
from fastapi.responses import StreamingResponse
from pydantic import BaseModel
from typing import Optional
from datetime import datetime, timezone, timedelta
from io import BytesIO
from pathlib import Path
import uuid

from database import db
from utils.auth import (
    get_current_user,
    get_current_user_from_token,
    require_permission,
    resolve_branch_filter,
)

router = APIRouter(prefix="/coach-notes", tags=["Coach Notes"])

PERMISSION_KEY = "coach-notes"
SAUDI_OFFSET = timedelta(hours=3)

# Preset rated categories (Arabic labels live in the report/UI layer).
CATEGORY_KEYS = ("punctuality", "member_handling", "performance")
CATEGORY_LABELS_AR = {
    "punctuality": "الالتزام بالمواعيد",
    "member_handling": "التعامل مع الأعضاء",
    "performance": "الأداء",
}


def _saudi_today() -> str:
    return (datetime.now(timezone.utc) + SAUDI_OFFSET).strftime("%Y-%m-%d")


def _branch_scope_filter(effective_branch: Optional[str]) -> dict:
    """Pin records to the caller's branch (admins → no restriction).

    Legacy/unscoped records (no branch_id) stay visible — same convention as
    coaches/coach_attendance routes.
    """
    if not effective_branch:
        return {}
    return {"$or": [
        {"branch_id": effective_branch},
        {"branch_id": None},
        {"branch_id": {"$exists": False}},
    ]}


def _clamp_rating(v) -> Optional[int]:
    if v is None or v == "":
        return None
    try:
        iv = int(v)
    except (TypeError, ValueError):
        return None
    return max(1, min(5, iv))


class NoteRatings(BaseModel):
    punctuality: Optional[int] = None
    member_handling: Optional[int] = None
    performance: Optional[int] = None


class CoachNoteCreate(BaseModel):
    coach_id: str
    date: Optional[str] = None
    note_text: Optional[str] = ""
    ratings: Optional[NoteRatings] = None


class CoachNoteUpdate(BaseModel):
    date: Optional[str] = None
    note_text: Optional[str] = None
    ratings: Optional[NoteRatings] = None


def _normalize_ratings(ratings: Optional[NoteRatings]) -> dict:
    data = ratings.model_dump() if ratings else {}
    return {k: _clamp_rating(data.get(k)) for k in CATEGORY_KEYS}


async def _resolve_coach_scoped(coach_id: str, current_user: dict) -> dict:
    """Fetch a coach the caller is allowed to see (branch-scoped)."""
    effective_branch = resolve_branch_filter(current_user, None)
    query = {"id": coach_id}
    if effective_branch:
        query.update(_branch_scope_filter(effective_branch))
    coach = await db.coaches.find_one(query, {"_id": 0})
    if not coach:
        raise HTTPException(status_code=404, detail="المدرب غير موجود")
    return coach


@router.post("")
async def create_coach_note(payload: CoachNoteCreate, current_user: dict = Depends(get_current_user)):
    await require_permission(current_user, PERMISSION_KEY)
    coach = await _resolve_coach_scoped(payload.coach_id, current_user)

    ratings = _normalize_ratings(payload.ratings)
    note_text = (payload.note_text or "").strip()
    if not note_text and all(v is None for v in ratings.values()):
        raise HTTPException(status_code=400, detail="اكتب ملاحظة أو قيّم تصنيف واحد على الأقل")

    now = datetime.now(timezone.utc).isoformat()
    doc = {
        "id": str(uuid.uuid4()),
        "coach_id": coach["id"],
        "coach_name": coach.get("name_ar") or coach.get("name") or "",
        "branch_id": coach.get("branch_id") or current_user.get("branch_id"),
        "date": (payload.date or _saudi_today())[:10],
        "note_text": note_text,
        "ratings": ratings,
        "created_by": current_user.get("user_id"),
        "created_by_name": current_user.get("username") or "",
        "created_at": now,
        "updated_at": now,
    }
    await db.coach_notes.insert_one(doc)
    doc.pop("_id", None)
    return doc


@router.get("")
async def list_coach_notes(
    coach_id: Optional[str] = None,
    month: Optional[str] = None,
    branch_filter: Optional[str] = None,
    current_user: dict = Depends(get_current_user),
):
    await require_permission(current_user, PERMISSION_KEY)
    effective_branch = resolve_branch_filter(current_user, branch_filter)

    query: dict = {}
    if effective_branch:
        query.update(_branch_scope_filter(effective_branch))
    if coach_id:
        query["coach_id"] = coach_id
    if month:
        query["date"] = {"$regex": f"^{month}"}

    notes = await db.coach_notes.find(query, {"_id": 0}).to_list(5000)
    notes.sort(key=lambda n: (n.get("date", ""), n.get("created_at", "")), reverse=True)
    return notes


@router.put("/{note_id}")
async def update_coach_note(note_id: str, payload: CoachNoteUpdate, current_user: dict = Depends(get_current_user)):
    await require_permission(current_user, PERMISSION_KEY)
    effective_branch = resolve_branch_filter(current_user, None)
    query = {"id": note_id}
    if effective_branch:
        query.update(_branch_scope_filter(effective_branch))

    update_data: dict = {}
    if payload.date is not None:
        update_data["date"] = payload.date[:10]
    if payload.note_text is not None:
        update_data["note_text"] = payload.note_text.strip()
    if payload.ratings is not None:
        update_data["ratings"] = _normalize_ratings(payload.ratings)
    if not update_data:
        raise HTTPException(status_code=400, detail="لا يوجد بيانات للتعديل")
    update_data["updated_at"] = datetime.now(timezone.utc).isoformat()

    result = await db.coach_notes.find_one_and_update(
        query, {"$set": update_data}, return_document=True
    )
    if not result:
        raise HTTPException(status_code=404, detail="الملاحظة غير موجودة")
    result.pop("_id", None)
    return result


@router.delete("/{note_id}")
async def delete_coach_note(note_id: str, current_user: dict = Depends(get_current_user)):
    await require_permission(current_user, PERMISSION_KEY)
    effective_branch = resolve_branch_filter(current_user, None)
    query = {"id": note_id}
    if effective_branch:
        query.update(_branch_scope_filter(effective_branch))
    result = await db.coach_notes.delete_one(query)
    if result.deleted_count == 0:
        raise HTTPException(status_code=404, detail="الملاحظة غير موجودة")
    return {"success": True}


def _averages(notes: list) -> dict:
    out = {}
    for key in CATEGORY_KEYS:
        vals = [n["ratings"].get(key) for n in notes
                if isinstance(n.get("ratings"), dict) and n["ratings"].get(key) is not None]
        out[key] = round(sum(vals) / len(vals), 1) if vals else None
    return out


async def _build_monthly_report(month: str, coach_id: Optional[str], current_user: dict) -> dict:
    effective_branch = resolve_branch_filter(current_user, None)
    query: dict = {"date": {"$regex": f"^{month}"}}
    if effective_branch:
        query.update(_branch_scope_filter(effective_branch))
    if coach_id:
        query["coach_id"] = coach_id

    notes = await db.coach_notes.find(query, {"_id": 0}).to_list(10000)

    grouped: dict = {}
    for n in notes:
        cid = n.get("coach_id") or ""
        grouped.setdefault(cid, {"coach_id": cid, "coach_name": n.get("coach_name") or "", "notes": []})
        grouped[cid]["notes"].append(n)

    coaches = []
    for cid, g in grouped.items():
        g["notes"].sort(key=lambda x: (x.get("date", ""), x.get("created_at", "")))
        coaches.append({
            "coach_id": cid,
            "coach_name": g["coach_name"],
            "count": len(g["notes"]),
            "averages": _averages(g["notes"]),
            "notes": g["notes"],
        })
    coaches.sort(key=lambda c: c["coach_name"])
    return {"month": month, "coaches": coaches, "total_notes": len(notes)}


@router.get("/monthly-report")
async def monthly_report(
    month: str,
    coach_id: Optional[str] = None,
    current_user: dict = Depends(get_current_user),
):
    await require_permission(current_user, PERMISSION_KEY)
    return await _build_monthly_report(month, coach_id, current_user)


def _get_openpyxl():
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
    return type('XL', (), {
        'Workbook': Workbook, 'Font': Font, 'Alignment': Alignment,
        'Border': Border, 'Side': Side, 'PatternFill': PatternFill,
    })()


def _get_reportlab():
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    return type('RL', (), {
        'colors': colors, 'A4': A4,
        'SimpleDocTemplate': SimpleDocTemplate, 'Table': Table,
        'TableStyle': TableStyle, 'Paragraph': Paragraph, 'Spacer': Spacer,
        'ParagraphStyle': ParagraphStyle,
        'mm': mm, 'pdfmetrics': pdfmetrics, 'TTFont': TTFont,
    })()


def _fmt_avg(v) -> str:
    return "-" if v is None else str(v)


def _fmt_rating(v) -> str:
    return "-" if v is None else str(v)


@router.get("/monthly-report/export")
async def export_monthly_report(
    month: str,
    coach_id: Optional[str] = None,
    format: str = "xlsx",
    token: Optional[str] = None,
    current_user: dict = Depends(get_current_user_from_token),
):
    await require_permission(current_user, PERMISSION_KEY)
    if format not in ("xlsx", "pdf"):
        raise HTTPException(status_code=400, detail="format يجب أن يكون xlsx أو pdf")

    report = await _build_monthly_report(month, coach_id, current_user)
    export_date = _saudi_today()
    headers_row = ["م", "المدرب", "التاريخ", "الالتزام بالمواعيد", "التعامل مع الأعضاء", "الأداء", "الملاحظة"]

    if format == "xlsx":
        xl = _get_openpyxl()
        Workbook = xl.Workbook; Font = xl.Font; PatternFill = xl.PatternFill
        Border = xl.Border; Side = xl.Side; Alignment = xl.Alignment
        wb = Workbook()
        ws = wb.active
        ws.title = "ملاحظات المدربين"

        orange_fill = PatternFill("solid", fgColor="F97316")
        white_on_orange = Font(bold=True, color="FFFFFF")
        border = Border(
            left=Side(style='thin', color='DDDDDD'), right=Side(style='thin', color='DDDDDD'),
            top=Side(style='thin', color='DDDDDD'), bottom=Side(style='thin', color='DDDDDD'),
        )
        center = Alignment(horizontal='center', vertical='center', wrap_text=True)
        right_align = Alignment(horizontal='right', vertical='center', wrap_text=True)

        col_count = len(headers_row)
        last_col_letter = chr(ord('A') + col_count - 1)
        ws.append(["شركة اداء الابطال العالمية للرياضة"])
        ws.append([f"التقرير الشهري لملاحظات المدربين — {month}"])
        ws.append([f"تاريخ التصدير: {export_date}"])
        ws.append([])
        for r in (1, 2, 3):
            ws.cell(row=r, column=1).font = Font(bold=True, size=13 if r == 1 else 11)
            ws.merge_cells(f'A{r}:{last_col_letter}{r}')
            ws[f'A{r}'].alignment = right_align

        header_row_idx = 5
        for col, h in enumerate(headers_row, 1):
            c = ws.cell(row=header_row_idx, column=col, value=h)
            c.fill = orange_fill; c.font = white_on_orange; c.border = border; c.alignment = center

        row_idx = header_row_idx + 1
        idx = 0
        for coach in report["coaches"]:
            avg = coach["averages"]
            for n in coach["notes"]:
                idx += 1
                ratings = n.get("ratings") or {}
                values = [
                    idx, coach["coach_name"], n.get("date", ""),
                    _fmt_rating(ratings.get("punctuality")),
                    _fmt_rating(ratings.get("member_handling")),
                    _fmt_rating(ratings.get("performance")),
                    n.get("note_text", ""),
                ]
                for col, v in enumerate(values, 1):
                    c = ws.cell(row=row_idx, column=col, value=v)
                    c.border = border
                    c.alignment = right_align if col in (2, 7) else center
                row_idx += 1
            # Coach average summary row
            summary = [
                "", f"متوسط {coach['coach_name']}", f"عدد الملاحظات: {coach['count']}",
                _fmt_avg(avg.get("punctuality")), _fmt_avg(avg.get("member_handling")),
                _fmt_avg(avg.get("performance")), "",
            ]
            for col, v in enumerate(summary, 1):
                c = ws.cell(row=row_idx, column=col, value=v)
                c.fill = PatternFill("solid", fgColor="FFF7ED")
                c.font = Font(bold=True)
                c.border = border
                c.alignment = right_align if col in (2, 3) else center
            row_idx += 1

        widths = [5, 22, 12, 16, 16, 10, 50]
        for col, w in enumerate(widths, 1):
            ws.column_dimensions[chr(ord('A') + col - 1)].width = w

        buffer = BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return StreamingResponse(
            buffer,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f"attachment; filename=coach_notes_{month}.xlsx"},
        )

    # ── PDF ──
    rl = _get_reportlab()
    colors = rl.colors; A4 = rl.A4; mm = rl.mm
    SimpleDocTemplate = rl.SimpleDocTemplate; Table = rl.Table
    TableStyle = rl.TableStyle; Paragraph = rl.Paragraph; Spacer = rl.Spacer
    ParagraphStyle = rl.ParagraphStyle
    pdfmetrics = rl.pdfmetrics; TTFont = rl.TTFont

    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, topMargin=15*mm, bottomMargin=15*mm,
                            leftMargin=12*mm, rightMargin=12*mm)
    font_name = "Helvetica"
    for fp in [
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
        "/usr/share/fonts/TTF/DejaVuSans.ttf",
        "/nix/store/dejavu-fonts/share/fonts/truetype/DejaVuSans.ttf",
    ]:
        try:
            if Path(fp).exists():
                pdfmetrics.registerFont(TTFont('ArabicFont', fp))
                font_name = 'ArabicFont'
                break
        except Exception:
            continue

    title_style = ParagraphStyle('T', fontName=font_name, fontSize=13, leading=18, alignment=2)
    sub_style = ParagraphStyle('S', fontName=font_name, fontSize=9, leading=13,
                               textColor=colors.HexColor('#555555'), alignment=2)
    coach_style = ParagraphStyle('CN', fontName=font_name, fontSize=11, leading=15,
                                 textColor=colors.HexColor('#9A3412'), alignment=2)
    cell_style = ParagraphStyle('C', fontName=font_name, fontSize=8, leading=11, alignment=1)
    note_cell_style = ParagraphStyle('NC', fontName=font_name, fontSize=8, leading=11, alignment=2)
    hdr_style = ParagraphStyle('H', fontName=font_name, fontSize=8, leading=11,
                               textColor=colors.white, alignment=1)

    elements = []
    elements.append(Paragraph("شركة اداء الابطال العالمية للرياضة", title_style))
    elements.append(Paragraph(f"التقرير الشهري لملاحظات المدربين — {month}", title_style))
    elements.append(Paragraph(f"تاريخ التصدير: {export_date}", sub_style))
    elements.append(Spacer(1, 5*mm))

    # RTL header order (rightmost first visually)
    pdf_headers = ["الملاحظة", "الأداء", "التعامل مع الأعضاء", "الالتزام بالمواعيد", "التاريخ", "م"]
    col_widths_pdf = [78*mm, 16*mm, 22*mm, 22*mm, 20*mm, 8*mm]

    if not report["coaches"]:
        elements.append(Paragraph("لا توجد ملاحظات في هذا الشهر", sub_style))

    for coach in report["coaches"]:
        avg = coach["averages"]
        elements.append(Spacer(1, 3*mm))
        elements.append(Paragraph(
            f"{coach['coach_name']} — عدد الملاحظات: {coach['count']} | "
            f"متوسط الالتزام: {_fmt_avg(avg.get('punctuality'))} | "
            f"التعامل: {_fmt_avg(avg.get('member_handling'))} | "
            f"الأداء: {_fmt_avg(avg.get('performance'))}",
            coach_style,
        ))
        data = [[Paragraph(h, hdr_style) for h in pdf_headers]]
        for i, n in enumerate(coach["notes"], 1):
            ratings = n.get("ratings") or {}
            data.append([
                Paragraph((n.get("note_text") or "").replace("\n", "<br/>") or "-", note_cell_style),
                Paragraph(_fmt_rating(ratings.get("performance")), cell_style),
                Paragraph(_fmt_rating(ratings.get("member_handling")), cell_style),
                Paragraph(_fmt_rating(ratings.get("punctuality")), cell_style),
                Paragraph(n.get("date", ""), cell_style),
                Paragraph(str(i), cell_style),
            ])
        table = Table(data, colWidths=col_widths_pdf, repeatRows=1)
        table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#F97316')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('GRID', (0, 0), (-1, -1), 0.4, colors.HexColor('#DDDDDD')),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#FFF7ED')]),
            ('TOPPADDING', (0, 0), (-1, -1), 4),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ]))
        elements.append(table)

    elements.append(Spacer(1, 5*mm))
    elements.append(Paragraph(f"إجمالي الملاحظات: {report['total_notes']}", sub_style))

    doc.build(elements)
    buffer.seek(0)
    return StreamingResponse(
        buffer,
        media_type="application/pdf",
        headers={"Content-Disposition": f"attachment; filename=coach_notes_{month}.pdf"},
    )
