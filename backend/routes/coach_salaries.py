"""Coach Salaries API - رواتب المدربين"""
from fastapi import APIRouter, HTTPException, Depends
from fastapi.responses import StreamingResponse
from pydantic import BaseModel
from typing import Optional, List, Dict, Any
from datetime import datetime, timezone, timedelta
from io import BytesIO
from pathlib import Path
import uuid

from .common import db, get_current_user
from utils.auth import require_branch_scope, resolve_branch_filter, require_permission

router = APIRouter(prefix="/coach-salaries", tags=["coach-salaries"])


def _branch_scope_filter(effective_branch: Optional[str]) -> dict:
    if not effective_branch:
        return {}
    return {"$or": [
        {"branch_id": effective_branch},
        {"branch_id": None},
        {"branch_id": {"$exists": False}}
    ]}


def _last_day_of_month(year_month: str) -> str:
    y, m = year_month.split("-")
    y, m = int(y), int(m)
    if m == 12:
        first_next = datetime(y + 1, 1, 1)
    else:
        first_next = datetime(y, m + 1, 1)
    last = first_next - timedelta(days=1)
    return last.strftime("%Y-%m-%d")


async def _next_expense_number() -> str:
    counter = await db.counters.find_one_and_update(
        {"_id": "internal_expense_number"},
        {"$inc": {"seq": 1}},
        upsert=True,
        return_document=True,
    )
    seq = (counter or {}).get("seq", 1)
    if seq < 10001:
        await db.counters.update_one(
            {"_id": "internal_expense_number"},
            {"$max": {"seq": 10001}}
        )
        seq = 10001
    return f"EXP-{seq:05d}"


async def _ensure_indexes():
    try:
        await db.coach_salaries.create_index(
            [("year_month", 1), ("coach_id", 1)],
            unique=True,
            name="uniq_year_month_coach"
        )
    except Exception:
        pass


def _compute_attendance_stats(coach: dict, attendance_records: List[dict]) -> Dict[str, Any]:
    threshold_str = coach.get("expected_checkin_time") or "09:00"
    try:
        threshold_dt = datetime.strptime(threshold_str, "%H:%M")
    except Exception:
        threshold_dt = datetime.strptime("09:00", "%H:%M")
        threshold_str = "09:00"

    try:
        monthly_work_days = max(1, min(int(coach.get("monthly_work_days") or 30), 31))
    except (ValueError, TypeError):
        monthly_work_days = 30
    contract_type = coach.get("contract_type") or "full_time"
    if contract_type not in ("full_time", "part_time"):
        contract_type = "full_time"

    coach_records = [r for r in attendance_records if r.get("coach_id") == coach["id"]]
    present_days = len([r for r in coach_records if r.get("status") in ("present", "checked_out")])
    leave_days = len([r for r in coach_records if r.get("status") == "leave"])
    absent_days = max(0, monthly_work_days - present_days - leave_days)

    late_minutes_total = 0
    for r in coach_records:
        if r.get("status") not in ("present", "checked_out"):
            continue
        cin_str = r.get("check_in_time")
        if not cin_str:
            continue
        try:
            cin_dt = datetime.strptime(cin_str, "%H:%M")
            diff = (cin_dt - threshold_dt).total_seconds() / 60
            if diff > 0:
                late_minutes_total += diff
        except Exception:
            pass

    return {
        "present_days": present_days,
        "absent_days": absent_days,
        "leave_days": leave_days,
        "monthly_work_days": monthly_work_days,
        "contract_type": contract_type,
        "late_minutes": int(round(late_minutes_total)),
        "threshold": threshold_str,
    }


@router.get("")
async def list_salaries(
    month: str,
    branch_filter: Optional[str] = None,
    current_user: dict = Depends(get_current_user)
):
    """Return computed salary rows for every coach in scope, merged with saved
    drafts/disbursed records for the given month."""
    await require_permission(current_user, "salaries")
    await _ensure_indexes()

    effective_branch = resolve_branch_filter(current_user, branch_filter)

    coach_query = _branch_scope_filter(effective_branch) if effective_branch else {}
    coaches = await db.coaches.find(coach_query, {"_id": 0}).to_list(500)

    att_query: dict = {"date": {"$regex": f"^{month}"}}
    if effective_branch:
        att_query.update(_branch_scope_filter(effective_branch))
    attendance = await db.coach_attendance.find(att_query, {"_id": 0}).to_list(20000)

    saved_query: dict = {"year_month": month}
    if effective_branch:
        saved_query.update(_branch_scope_filter(effective_branch))
    saved = await db.coach_salaries.find(saved_query, {"_id": 0}).to_list(2000)
    saved_by_coach = {s["coach_id"]: s for s in saved}

    adv_query: dict = {
        "advance_date": {"$lte": _last_day_of_month(month)},
    }
    if effective_branch:
        adv_query.update(_branch_scope_filter(effective_branch))
    advances = await db.coach_advances.find(adv_query, {"_id": 0}).sort("advance_date", 1).to_list(5000)

    rows: List[dict] = []
    for coach in coaches:
        cid = coach["id"]
        base_salary = float(coach.get("base_salary") or 0)
        daily_rate = float(coach.get("daily_deduction_rate") or 0)
        late_rate = float(coach.get("late_minute_rate") or 0)
        stats = _compute_attendance_stats(coach, attendance)
        contract_type = stats["contract_type"]
        monthly_work_days = stats["monthly_work_days"]
        deduction_absent = round(stats["absent_days"] * daily_rate, 2)
        deduction_late = round(stats["late_minutes"] * late_rate, 2)

        coach_advances = [
            a for a in advances
            if a.get("coach_id") == cid and a.get("status") == "pending"
        ]

        common_fields = {
            "year_month": month,
            "coach_id": cid,
            "coach_name": coach.get("name_ar") or coach.get("name", ""),
            "branch_id": coach.get("branch_id"),
            "base_salary": base_salary,
            "daily_deduction_rate": daily_rate,
            "late_minute_rate": late_rate,
            "contract_type": contract_type,
            "monthly_work_days": monthly_work_days,
            "present_days": stats["present_days"],
            "absent_days": stats["absent_days"],
            "leave_days": stats["leave_days"],
            "late_minutes": stats["late_minutes"],
            "deduction_absent": deduction_absent,
            "deduction_late": deduction_late,
            "available_advances": coach_advances,
        }

        existing = saved_by_coach.get(cid)
        if existing:
            advances_repaid_ids = existing.get("advances_repaid", []) or []
            advances_repaid_total = float(existing.get("advances_repaid_total") or 0)
            bonus = float(existing.get("bonus") or 0)
            manual_deductions = existing.get("manual_deductions") or []
            manual_total = sum(float(m.get("amount") or 0) for m in manual_deductions)
            net = round(
                base_salary - deduction_absent - deduction_late
                + bonus - manual_total - advances_repaid_total,
                2
            )
            rows.append({
                **common_fields,
                "id": existing.get("id"),
                "bonus": bonus,
                "manual_deductions": manual_deductions,
                "advances_repaid": advances_repaid_ids,
                "advances_repaid_total": advances_repaid_total,
                "net_amount": existing.get("net_amount", net) if existing.get("status") == "disbursed" else net,
                "status": existing.get("status", "draft"),
                "expense_id": existing.get("expense_id"),
                "disbursed_at": existing.get("disbursed_at"),
                "disbursed_by": existing.get("disbursed_by"),
                "notes": existing.get("notes", ""),
            })
        else:
            advances_total = sum(float(a.get("amount") or 0) for a in coach_advances)
            net = round(
                base_salary - deduction_absent - deduction_late - advances_total,
                2
            )
            rows.append({
                **common_fields,
                "id": None,
                "bonus": 0,
                "manual_deductions": [],
                "advances_repaid": [a["id"] for a in coach_advances],
                "advances_repaid_total": advances_total,
                "net_amount": net,
                "status": "new",
                "expense_id": None,
                "disbursed_at": None,
                "disbursed_by": None,
                "notes": "",
            })

    rows.sort(key=lambda r: r["coach_name"] or "")
    return {"month": month, "rows": rows}


class ManualDeduction(BaseModel):
    description: str
    amount: float


class SalarySave(BaseModel):
    year_month: str
    coach_id: str
    bonus: Optional[float] = 0
    manual_deductions: List[ManualDeduction] = []
    advances_repaid: List[str] = []
    notes: Optional[str] = ""


@router.post("/save")
async def save_salary_draft(
    data: SalarySave,
    current_user: dict = Depends(get_current_user)
):
    await require_permission(current_user, "salaries")
    await _ensure_indexes()
    effective_branch = require_branch_scope(current_user)

    coach = await db.coaches.find_one({"id": data.coach_id}, {"_id": 0})
    if not coach:
        raise HTTPException(status_code=404, detail="المدرب غير موجود")
    if effective_branch:
        coach_branch = coach.get("branch_id")
        if coach_branch and coach_branch != effective_branch:
            raise HTTPException(status_code=404, detail="المدرب غير موجود")

    existing = await db.coach_salaries.find_one(
        {"year_month": data.year_month, "coach_id": data.coach_id},
        {"_id": 0}
    )
    if existing and existing.get("status") == "disbursed":
        raise HTTPException(status_code=400, detail="هذا الراتب مصروف بالفعل ولا يمكن تعديله. ألغِ الصرف أولاً.")

    base_salary = float(coach.get("base_salary") or 0)
    daily_rate = float(coach.get("daily_deduction_rate") or 0)
    late_rate = float(coach.get("late_minute_rate") or 0)

    att_query = {"date": {"$regex": f"^{data.year_month}"}, "coach_id": data.coach_id}
    attendance = await db.coach_attendance.find(att_query, {"_id": 0}).to_list(5000)
    stats = _compute_attendance_stats(coach, attendance)
    deduction_absent = round(stats["absent_days"] * daily_rate, 2)
    deduction_late = round(stats["late_minutes"] * late_rate, 2)

    advances_total = 0.0
    if data.advances_repaid:
        existing_id = existing.get("id") if existing else None
        adv_docs = await db.coach_advances.find(
            {"id": {"$in": data.advances_repaid}, "coach_id": data.coach_id},
            {"_id": 0}
        ).to_list(500)
        if len(adv_docs) != len(set(data.advances_repaid)):
            raise HTTPException(status_code=400, detail="إحدى السُلف المختارة غير موجودة أو لا تخص هذا المدرب")
        for a in adv_docs:
            if a.get("status") == "repaid" and a.get("repaid_in_salary_id") != existing_id:
                raise HTTPException(status_code=400, detail=f"السلفة بتاريخ {a.get('advance_date')} مخصومة من راتب آخر")
            advances_total += float(a.get("amount") or 0)

    bonus = float(data.bonus or 0)
    manual_total = sum(float(m.amount) for m in data.manual_deductions)
    net = round(
        base_salary - deduction_absent - deduction_late
        + bonus - manual_total - advances_total,
        2
    )

    now = datetime.now(timezone.utc).isoformat()
    if existing:
        salary_id = existing["id"]
        await db.coach_salaries.update_one(
            {"id": salary_id},
            {"$set": {
                "base_salary": base_salary,
                "daily_deduction_rate": daily_rate,
                "late_minute_rate": late_rate,
                "contract_type": stats["contract_type"],
                "monthly_work_days": stats["monthly_work_days"],
                "present_days": stats["present_days"],
                "absent_days": stats["absent_days"],
                "leave_days": stats["leave_days"],
                "late_minutes": stats["late_minutes"],
                "deduction_absent": deduction_absent,
                "deduction_late": deduction_late,
                "bonus": bonus,
                "manual_deductions": [m.dict() for m in data.manual_deductions],
                "advances_repaid": data.advances_repaid,
                "advances_repaid_total": round(advances_total, 2),
                "net_amount": net,
                "notes": data.notes or "",
                "updated_at": now,
            }}
        )
    else:
        salary_id = str(uuid.uuid4())
        doc = {
            "id": salary_id,
            "year_month": data.year_month,
            "coach_id": data.coach_id,
            "coach_name": coach.get("name_ar") or coach.get("name", ""),
            "branch_id": coach.get("branch_id"),
            "base_salary": base_salary,
            "daily_deduction_rate": daily_rate,
            "late_minute_rate": late_rate,
            "contract_type": stats["contract_type"],
            "monthly_work_days": stats["monthly_work_days"],
            "present_days": stats["present_days"],
            "absent_days": stats["absent_days"],
            "leave_days": stats["leave_days"],
            "late_minutes": stats["late_minutes"],
            "deduction_absent": deduction_absent,
            "deduction_late": deduction_late,
            "bonus": bonus,
            "manual_deductions": [m.dict() for m in data.manual_deductions],
            "advances_repaid": data.advances_repaid,
            "advances_repaid_total": round(advances_total, 2),
            "net_amount": net,
            "status": "draft",
            "expense_id": None,
            "disbursed_at": None,
            "disbursed_by": None,
            "notes": data.notes or "",
            "created_by": current_user.get("username") or current_user.get("name", ""),
            "created_at": now,
            "updated_at": now,
        }
        try:
            await db.coach_salaries.insert_one(doc)
        except Exception as e:
            if "duplicate key" in str(e).lower() or "E11000" in str(e):
                raise HTTPException(status_code=400, detail="يوجد سجل راتب لهذا المدرب لنفس الشهر بالفعل")
            raise

    saved = await db.coach_salaries.find_one({"id": salary_id}, {"_id": 0})
    return saved


@router.post("/{salary_id}/disburse")
async def disburse_salary(
    salary_id: str,
    current_user: dict = Depends(get_current_user)
):
    await require_permission(current_user, "salaries")
    effective_branch = resolve_branch_filter(current_user, None)
    query: dict = {"id": salary_id}
    if effective_branch:
        query.update(_branch_scope_filter(effective_branch))

    salary = await db.coach_salaries.find_one(query, {"_id": 0})
    if not salary:
        raise HTTPException(status_code=404, detail="سجل الراتب غير موجود")
    if salary.get("status") == "disbursed":
        raise HTTPException(status_code=400, detail="الراتب مصروف بالفعل")

    net = float(salary.get("net_amount") or 0)
    if net <= 0:
        raise HTTPException(status_code=400, detail="صافي الراتب صفر أو سالب — لا يمكن الصرف")

    now_iso = datetime.now(timezone.utc).isoformat()
    expense_id = str(uuid.uuid4())

    # Compare-and-set salary status to lock against concurrent disburse
    lock = await db.coach_salaries.update_one(
        {"id": salary_id, "status": {"$ne": "disbursed"}},
        {"$set": {
            "status": "disbursed",
            "expense_id": expense_id,
            "disbursed_at": now_iso,
            "disbursed_by": current_user.get("username") or current_user.get("name", ""),
            "updated_at": now_iso,
        }}
    )
    if lock.modified_count == 0:
        raise HTTPException(status_code=409, detail="تم صرف الراتب بالفعل من جلسة أخرى")

    advances_repaid = salary.get("advances_repaid") or []
    if advances_repaid:
        adv_result = await db.coach_advances.update_many(
            {
                "id": {"$in": advances_repaid},
                "coach_id": salary.get("coach_id"),
                "status": "pending",
            },
            {"$set": {
                "status": "repaid",
                "repaid_in_salary_id": salary_id,
                "repaid_at": now_iso,
            }}
        )
        if adv_result.modified_count != len(advances_repaid):
            # Roll back salary status — at least one advance was already repaid/missing
            await db.coach_salaries.update_one(
                {"id": salary_id},
                {"$set": {
                    "status": "draft" if salary.get("status") == "draft" else salary.get("status", "draft"),
                    "expense_id": None,
                    "disbursed_at": None,
                    "disbursed_by": None,
                    "updated_at": now_iso,
                }}
            )
            await db.coach_advances.update_many(
                {"repaid_in_salary_id": salary_id, "repaid_at": now_iso},
                {"$set": {
                    "status": "pending",
                    "repaid_in_salary_id": None,
                    "repaid_at": None,
                }}
            )
            raise HTTPException(status_code=409, detail="إحدى السُلف لم تعد متاحة للخصم — تم إلغاء العملية")

    expense_number = await _next_expense_number()
    expense_date = _last_day_of_month(salary["year_month"])
    expense_doc = {
        "id": expense_id,
        "expense_number": expense_number,
        "expense_date": expense_date,
        "expense_type": "coach_salary",
        "cost_center": salary.get("branch_id"),
        "description": f"راتب {salary.get('coach_name', '')} عن شهر {salary['year_month']}",
        "amount": float(net),
        "payment_method": "cash",
        "executor_name": salary.get("coach_name", ""),
        "notes": salary.get("notes", ""),
        "receipt_url": None,
        "status": "approved",
        "branch_id": salary.get("branch_id"),
        "created_by": current_user.get("username") or current_user.get("name", ""),
        "created_at": now_iso,
        "journal_entry_id": None,
        "coach_id": salary.get("coach_id"),
        "coach_salary_id": salary_id,
    }
    try:
        await db.internal_expenses.insert_one(expense_doc)
    except Exception as exc:
        await db.coach_salaries.update_one(
            {"id": salary_id},
            {"$set": {
                "status": "draft",
                "expense_id": None,
                "disbursed_at": None,
                "disbursed_by": None,
                "updated_at": now_iso,
            }}
        )
        if advances_repaid:
            await db.coach_advances.update_many(
                {"repaid_in_salary_id": salary_id, "repaid_at": now_iso},
                {"$set": {
                    "status": "pending",
                    "repaid_in_salary_id": None,
                    "repaid_at": None,
                }}
            )
        raise HTTPException(status_code=500, detail=f"فشل إنشاء سجل المصروف — تم التراجع: {exc}")

    updated = await db.coach_salaries.find_one({"id": salary_id}, {"_id": 0})

    try:
        from .notifications import send_to_coach
        await send_to_coach(
            coach_id=salary.get("coach_id"),
            title="تم صرف الراتب",
            message=f"تم صرف راتب شهر {salary['year_month']} بقيمة {float(net):,.2f} ريال",
            notif_type="salary_disbursed",
            link=f"/coach-salaries/{salary_id}/payslip.pdf",
            branch_id=salary.get("branch_id"),
            tag=f"salary-{salary_id}",
        )
    except Exception:
        import logging
        logging.getLogger(__name__).exception("Failed to notify coach of salary disbursement")

    return updated


@router.get("/compute")
async def compute_salaries_alias(
    month: str,
    branch_filter: Optional[str] = None,
    current_user: dict = Depends(get_current_user)
):
    return await get_salaries(month=month, branch_filter=branch_filter, current_user=current_user)


@router.post("/{salary_id}/cancel-disburse")
async def cancel_disburse(
    salary_id: str,
    current_user: dict = Depends(get_current_user)
):
    await require_permission(current_user, "salaries")
    if not current_user.get("is_admin", False):
        raise HTTPException(status_code=403, detail="صلاحية الأدمن مطلوبة لإلغاء الصرف")

    now_iso = datetime.now(timezone.utc).isoformat()
    lock = await db.coach_salaries.find_one_and_update(
        {"id": salary_id, "status": "disbursed"},
        {"$set": {
            "status": "draft",
            "expense_id": None,
            "disbursed_at": None,
            "disbursed_by": None,
            "updated_at": now_iso,
        }},
    )
    if not lock:
        raise HTTPException(status_code=400, detail="الراتب غير مصروف أو تم إلغاؤه من جلسة أخرى")

    salary = lock
    if salary.get("expense_id"):
        await db.internal_expenses.delete_one({"id": salary["expense_id"]})

    advances_repaid = salary.get("advances_repaid") or []
    if advances_repaid:
        await db.coach_advances.update_many(
            {"id": {"$in": advances_repaid}, "repaid_in_salary_id": salary_id},
            {"$set": {
                "status": "pending",
                "repaid_in_salary_id": None,
                "repaid_at": None,
            }}
        )

    updated = await db.coach_salaries.find_one({"id": salary_id}, {"_id": 0})
    return updated


@router.delete("/{salary_id}")
async def delete_salary(
    salary_id: str,
    current_user: dict = Depends(get_current_user)
):
    await require_permission(current_user, "salaries")
    effective_branch = resolve_branch_filter(current_user, None)
    query: dict = {"id": salary_id}
    if effective_branch:
        query.update(_branch_scope_filter(effective_branch))

    salary = await db.coach_salaries.find_one(query, {"_id": 0})
    if not salary:
        raise HTTPException(status_code=404, detail="سجل الراتب غير موجود")
    if salary.get("status") == "disbursed":
        raise HTTPException(status_code=400, detail="لا يمكن حذف راتب مصروف. ألغِ الصرف أولاً")

    await db.coach_salaries.delete_one({"id": salary_id})
    return {"message": "تم الحذف"}


class BulkSalaryItem(BaseModel):
    coach_id: str
    bonus: Optional[float] = 0
    manual_deductions: List[ManualDeduction] = []
    advances_repaid: List[str] = []
    notes: Optional[str] = ""


class BulkSaveRequest(BaseModel):
    year_month: str
    items: List[BulkSalaryItem]


@router.post("/bulk-save")
async def bulk_save_drafts(
    data: BulkSaveRequest,
    current_user: dict = Depends(get_current_user)
):
    """حفظ الكل كمسودة - save drafts for many coaches at once."""
    await require_permission(current_user, "salaries")
    saved = []
    errors = []
    for item in data.items:
        try:
            payload = SalarySave(
                year_month=data.year_month,
                coach_id=item.coach_id,
                bonus=item.bonus,
                manual_deductions=item.manual_deductions,
                advances_repaid=item.advances_repaid,
                notes=item.notes,
            )
            res = await save_salary_draft(payload, current_user)
            saved.append(res)
        except HTTPException as e:
            errors.append({"coach_id": item.coach_id, "error": e.detail})
        except Exception as e:
            errors.append({"coach_id": item.coach_id, "error": str(e)})
    return {"saved": saved, "errors": errors, "saved_count": len(saved), "error_count": len(errors)}


class BulkDisburseRequest(BaseModel):
    salary_ids: List[str]


@router.post("/bulk-disburse")
async def bulk_disburse(
    data: BulkDisburseRequest,
    current_user: dict = Depends(get_current_user)
):
    """صرف الكل - disburse multiple salaries at once."""
    await require_permission(current_user, "salaries")
    disbursed = []
    errors = []
    for sid in data.salary_ids:
        try:
            res = await disburse_salary(sid, current_user)
            disbursed.append(res)
        except HTTPException as e:
            errors.append({"salary_id": sid, "error": e.detail})
        except Exception as e:
            errors.append({"salary_id": sid, "error": str(e)})
    return {"disbursed": disbursed, "errors": errors, "disbursed_count": len(disbursed), "error_count": len(errors)}


def _months_between(from_ym: str, to_ym: str) -> List[str]:
    import re
    if not (isinstance(from_ym, str) and isinstance(to_ym, str)
            and re.fullmatch(r"\d{4}-\d{2}", from_ym)
            and re.fullmatch(r"\d{4}-\d{2}", to_ym)):
        raise HTTPException(status_code=400, detail="from/to يجب أن يكون بصيغة YYYY-MM")
    fy, fm = (int(x) for x in from_ym.split("-"))
    ty, tm = (int(x) for x in to_ym.split("-"))
    if not (1 <= fm <= 12 and 1 <= tm <= 12):
        raise HTTPException(status_code=400, detail="رقم الشهر يجب أن يكون بين 1 و 12")
    if not (1900 <= fy <= 2999 and 1900 <= ty <= 2999):
        raise HTTPException(status_code=400, detail="السنة خارج النطاق المسموح")
    start = fy * 12 + (fm - 1)
    end = ty * 12 + (tm - 1)
    if end < start:
        raise HTTPException(status_code=400, detail="نطاق التاريخ غير صحيح: 'إلى' قبل 'من'")
    if (end - start) > 35:
        raise HTTPException(status_code=400, detail="النطاق كبير جداً (الحد الأقصى 36 شهراً)")
    out: List[str] = []
    for k in range(start, end + 1):
        y, m0 = divmod(k, 12)
        out.append(f"{y:04d}-{m0 + 1:02d}")
    return out


@router.get("/report")
async def salaries_report(
    from_month: str,
    to_month: str,
    branch_filter: Optional[str] = None,
    format: str = "json",
    current_user: dict = Depends(get_current_user)
):
    """تقرير شامل عبر عدة أشهر — per-coach aggregates + per-month history.

    Returns JSON by default; ``format=xlsx`` / ``format=pdf`` stream a file.
    """
    await require_permission(current_user, "salaries")
    if format not in ("json", "xlsx", "pdf"):
        raise HTTPException(status_code=400, detail="format يجب أن يكون json أو xlsx أو pdf")

    months = _months_between(from_month, to_month)
    effective_branch = resolve_branch_filter(current_user, branch_filter)

    coach_query = _branch_scope_filter(effective_branch) if effective_branch else {}
    coaches = await db.coaches.find(coach_query, {"_id": 0}).to_list(500)
    coaches_by_id = {c["id"]: c for c in coaches}

    sal_query: dict = {"year_month": {"$in": months}}
    if effective_branch:
        sal_query.update(_branch_scope_filter(effective_branch))
    salaries = await db.coach_salaries.find(sal_query, {"_id": 0}).to_list(20000)

    pending_query: dict = {"status": "pending"}
    if effective_branch:
        pending_query.update(_branch_scope_filter(effective_branch))
    pending_advances = await db.coach_advances.find(pending_query, {"_id": 0}).to_list(10000)
    pending_by_coach: Dict[str, float] = {}
    pending_count_by_coach: Dict[str, int] = {}
    for a in pending_advances:
        cid = a.get("coach_id")
        pending_by_coach[cid] = pending_by_coach.get(cid, 0.0) + float(a.get("amount") or 0)
        pending_count_by_coach[cid] = pending_count_by_coach.get(cid, 0) + 1

    rows_by_coach: Dict[str, Dict[str, Any]] = {}
    for s in salaries:
        cid = s.get("coach_id")
        if cid not in coaches_by_id:
            continue
        bucket = rows_by_coach.setdefault(cid, {
            "coach_id": cid,
            "coach_name": s.get("coach_name") or coaches_by_id[cid].get("name_ar") or coaches_by_id[cid].get("name", ""),
            "branch_id": s.get("branch_id") or coaches_by_id[cid].get("branch_id"),
            "total_base": 0.0,
            "total_deduction_absent": 0.0,
            "total_deduction_late": 0.0,
            "total_manual_deductions": 0.0,
            "total_bonus": 0.0,
            "total_advances_repaid": 0.0,
            "total_net_disbursed": 0.0,
            "total_net_all": 0.0,
            "absent_days_total": 0,
            "present_days_total": 0,
            "leave_days_total": 0,
            "late_minutes_total": 0,
            "months_recorded": 0,
            "months_disbursed": 0,
            "history": [],
        })
        manual_total = sum(float(m.get("amount") or 0) for m in (s.get("manual_deductions") or []))
        net = float(s.get("net_amount") or 0)
        is_disbursed = s.get("status") == "disbursed"
        bucket["total_base"] += float(s.get("base_salary") or 0)
        bucket["total_deduction_absent"] += float(s.get("deduction_absent") or 0)
        bucket["total_deduction_late"] += float(s.get("deduction_late") or 0)
        bucket["total_manual_deductions"] += manual_total
        bucket["total_bonus"] += float(s.get("bonus") or 0)
        bucket["total_advances_repaid"] += float(s.get("advances_repaid_total") or 0)
        bucket["total_net_all"] += net
        if is_disbursed:
            bucket["total_net_disbursed"] += net
            bucket["months_disbursed"] += 1
        bucket["absent_days_total"] += int(s.get("absent_days") or 0)
        bucket["present_days_total"] += int(s.get("present_days") or 0)
        bucket["leave_days_total"] += int(s.get("leave_days") or 0)
        bucket["late_minutes_total"] += int(s.get("late_minutes") or 0)
        bucket["months_recorded"] += 1
        bucket["history"].append({
            "year_month": s.get("year_month"),
            "base_salary": float(s.get("base_salary") or 0),
            "deduction_absent": float(s.get("deduction_absent") or 0),
            "deduction_late": float(s.get("deduction_late") or 0),
            "manual_deductions_total": manual_total,
            "bonus": float(s.get("bonus") or 0),
            "advances_repaid_total": float(s.get("advances_repaid_total") or 0),
            "absent_days": int(s.get("absent_days") or 0),
            "present_days": int(s.get("present_days") or 0),
            "late_minutes": int(s.get("late_minutes") or 0),
            "net_amount": net,
            "status": s.get("status"),
            "disbursed_at": s.get("disbursed_at"),
        })

    # ensure every in-scope coach appears even if no records
    for c in coaches:
        cid = c["id"]
        if cid not in rows_by_coach:
            rows_by_coach[cid] = {
                "coach_id": cid,
                "coach_name": c.get("name_ar") or c.get("name", ""),
                "branch_id": c.get("branch_id"),
                "total_base": 0.0, "total_deduction_absent": 0.0, "total_deduction_late": 0.0,
                "total_manual_deductions": 0.0, "total_bonus": 0.0, "total_advances_repaid": 0.0,
                "total_net_disbursed": 0.0, "total_net_all": 0.0,
                "absent_days_total": 0, "present_days_total": 0, "leave_days_total": 0,
                "late_minutes_total": 0, "months_recorded": 0, "months_disbursed": 0,
                "history": [],
            }

    coach_rows: List[Dict[str, Any]] = []
    for cid, b in rows_by_coach.items():
        b["history"].sort(key=lambda h: h.get("year_month") or "")
        b["pending_advances_total"] = round(pending_by_coach.get(cid, 0.0), 2)
        b["pending_advances_count"] = pending_count_by_coach.get(cid, 0)
        b["total_deductions"] = round(
            b["total_deduction_absent"] + b["total_deduction_late"] + b["total_manual_deductions"], 2
        )
        b["late_minutes_avg"] = round(
            b["late_minutes_total"] / b["months_recorded"], 1
        ) if b["months_recorded"] else 0.0
        for k in ("total_base", "total_deduction_absent", "total_deduction_late",
                  "total_manual_deductions", "total_bonus", "total_advances_repaid",
                  "total_net_disbursed", "total_net_all"):
            b[k] = round(b[k], 2)
        coach_rows.append(b)

    coach_rows.sort(key=lambda r: r["coach_name"] or "")

    totals = {
        "total_base": round(sum(r["total_base"] for r in coach_rows), 2),
        "total_deductions": round(sum(r["total_deductions"] for r in coach_rows), 2),
        "total_advances_repaid": round(sum(r["total_advances_repaid"] for r in coach_rows), 2),
        "total_net_disbursed": round(sum(r["total_net_disbursed"] for r in coach_rows), 2),
        "total_net_all": round(sum(r["total_net_all"] for r in coach_rows), 2),
        "pending_advances_total": round(sum(r["pending_advances_total"] for r in coach_rows), 2),
        "absent_days_total": sum(r["absent_days_total"] for r in coach_rows),
        "late_minutes_total": sum(r["late_minutes_total"] for r in coach_rows),
        "coaches_count": len(coach_rows),
        "months_count": len(months),
    }

    if format == "json":
        return {
            "from": from_month,
            "to": to_month,
            "months": months,
            "coaches": coach_rows,
            "totals": totals,
        }

    export_date = datetime.now().strftime("%Y-%m-%d")
    period_label = f"{from_month} → {to_month}"

    if format == "xlsx":
        from openpyxl import Workbook
        from openpyxl.styles import Font, Alignment, Border, Side, PatternFill

        wb = Workbook()
        ws = wb.active
        ws.title = "ملخص المدربين"

        orange_fill = PatternFill("solid", fgColor="F97316")
        white_on_orange = Font(bold=True, color="FFFFFF")
        alt_fill = PatternFill("solid", fgColor="FFF7ED")
        thin = Side(style='thin', color='DDDDDD')
        border = Border(left=thin, right=thin, top=thin, bottom=thin)
        center = Alignment(horizontal='center', vertical='center', wrap_text=True)
        right_align = Alignment(horizontal='right', vertical='center', wrap_text=True)

        ws.append(["شركة اداء الابطال العالمية للرياضة"])
        ws.append([f"تقرير رواتب وسُلف المدربين — {period_label}"])
        ws.append([f"تاريخ التصدير: {export_date}"])
        ws.append([])

        headers_row = [
            "م", "المدرب", "أشهر مسجّلة", "أشهر مصروفة",
            "إجمالي الراتب الأساسي", "إجمالي الخصومات",
            "إجمالي السُلف المخصومة", "إجمالي المصروف",
            "السُلف المعلّقة", "أيام الغياب", "متوسط دقائق التأخر",
        ]
        for col_idx in range(1, len(headers_row) + 1):
            ws.cell(row=1, column=col_idx).font = Font(bold=True, size=13)
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers_row))
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(headers_row))
        ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=len(headers_row))
        ws['A1'].alignment = right_align
        ws['A2'].alignment = right_align
        ws['A3'].alignment = right_align

        ws.append(headers_row)
        header_row_num = 5
        for col_idx in range(1, len(headers_row) + 1):
            cell = ws.cell(row=header_row_num, column=col_idx)
            cell.fill = orange_fill
            cell.font = white_on_orange
            cell.alignment = center
            cell.border = border

        for i, r in enumerate(coach_rows, 1):
            ws.append([
                i,
                r["coach_name"],
                r["months_recorded"],
                r["months_disbursed"],
                r["total_base"],
                r["total_deductions"],
                r["total_advances_repaid"],
                r["total_net_disbursed"],
                r["pending_advances_total"],
                r["absent_days_total"],
                r["late_minutes_avg"],
            ])
            row_num = header_row_num + i
            for col_idx in range(1, len(headers_row) + 1):
                cell = ws.cell(row=row_num, column=col_idx)
                if i % 2 == 0:
                    cell.fill = alt_fill
                cell.border = border
                cell.alignment = center if col_idx != 2 else right_align

        # Totals row
        totals_row = [
            "", "الإجمالي", "", "",
            totals["total_base"], totals["total_deductions"],
            totals["total_advances_repaid"], totals["total_net_disbursed"],
            totals["pending_advances_total"], totals["absent_days_total"], "",
        ]
        ws.append(totals_row)
        tot_row_num = header_row_num + len(coach_rows) + 1
        bold = Font(bold=True)
        for col_idx in range(1, len(headers_row) + 1):
            cell = ws.cell(row=tot_row_num, column=col_idx)
            cell.font = bold
            cell.fill = PatternFill("solid", fgColor="FEF3C7")
            cell.border = border
            cell.alignment = center if col_idx != 2 else right_align

        widths = [5, 26, 12, 14, 18, 16, 18, 16, 16, 12, 18]
        for idx, w in enumerate(widths, 1):
            ws.column_dimensions[ws.cell(row=1, column=idx).column_letter].width = w

        # Per-coach history sheets — one combined sheet with sections
        ws2 = wb.create_sheet("سجل تاريخي")
        ws2.append([f"السجل التاريخي للرواتب — {period_label}"])
        ws2.merge_cells(start_row=1, start_column=1, end_row=1, end_column=8)
        ws2['A1'].font = Font(bold=True, size=13)
        ws2['A1'].alignment = right_align
        ws2.append([])

        hist_headers = ["الشهر", "الراتب", "خصم غياب", "خصم تأخر", "علاوة", "سُلف مخصومة", "الصافي", "الحالة"]
        for r in coach_rows:
            ws2.append([r["coach_name"]])
            name_row = ws2.max_row
            ws2.cell(row=name_row, column=1).font = Font(bold=True, color="FFFFFF")
            ws2.cell(row=name_row, column=1).fill = orange_fill
            ws2.merge_cells(start_row=name_row, start_column=1, end_row=name_row, end_column=8)
            ws2.cell(row=name_row, column=1).alignment = right_align

            ws2.append(hist_headers)
            hdr = ws2.max_row
            for col_idx in range(1, 9):
                c = ws2.cell(row=hdr, column=col_idx)
                c.fill = PatternFill("solid", fgColor="FFF7ED")
                c.font = Font(bold=True)
                c.border = border
                c.alignment = center

            if not r["history"]:
                ws2.append(["لا توجد سجلات في هذه الفترة"])
                ws2.merge_cells(start_row=ws2.max_row, start_column=1, end_row=ws2.max_row, end_column=8)
                ws2.cell(row=ws2.max_row, column=1).alignment = center
            else:
                for h in r["history"]:
                    ws2.append([
                        h["year_month"],
                        h["base_salary"],
                        h["deduction_absent"],
                        h["deduction_late"],
                        h["bonus"],
                        h["advances_repaid_total"],
                        h["net_amount"],
                        "مصروف" if h["status"] == "disbursed" else ("مسودة" if h["status"] == "draft" else "—"),
                    ])
                    rn = ws2.max_row
                    for col_idx in range(1, 9):
                        ws2.cell(row=rn, column=col_idx).border = border
                        ws2.cell(row=rn, column=col_idx).alignment = center
            ws2.append([])

        for idx, w in enumerate([14, 14, 14, 14, 14, 14, 14, 12], 1):
            ws2.column_dimensions[ws2.cell(row=1, column=idx).column_letter].width = w

        buffer = BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        filename = f"coach_salaries_report_{from_month}_{to_month}.xlsx"
        return StreamingResponse(
            buffer,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f"attachment; filename={filename}"}
        )

    # PDF
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    font_name = "Helvetica"
    for fp in [
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
        "/usr/share/fonts/TTF/DejaVuSans.ttf",
        "/nix/store/dejavu-fonts/share/fonts/truetype/DejaVuSans.ttf",
    ]:
        try:
            if Path(fp).exists():
                pdfmetrics.registerFont(TTFont('ArabicFont', fp))
                font_name = 'ArabicFont'
                break
        except Exception:
            continue

    title_style = ParagraphStyle('T', fontName=font_name, fontSize=13, leading=18, alignment=2)
    sub_style = ParagraphStyle('S', fontName=font_name, fontSize=9, leading=13,
                               textColor=colors.HexColor('#555555'), alignment=2)
    cell_style = ParagraphStyle('C', fontName=font_name, fontSize=8, leading=11, alignment=1)
    hdr_style = ParagraphStyle('H', fontName=font_name, fontSize=8, leading=11,
                               textColor=colors.white, alignment=1)

    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=landscape(A4),
                            topMargin=12*mm, bottomMargin=12*mm,
                            leftMargin=10*mm, rightMargin=10*mm)
    elements = []
    elements.append(Paragraph("شركة اداء الابطال العالمية للرياضة", title_style))
    elements.append(Paragraph(f"تقرير رواتب وسُلف المدربين — {period_label}", title_style))
    elements.append(Paragraph(f"تاريخ التصدير: {export_date} — عدد المدربين: {totals['coaches_count']}", sub_style))
    elements.append(Spacer(1, 4*mm))

    pdf_headers = [
        "متوسط دقائق التأخر", "أيام الغياب", "السُلف المعلّقة",
        "إجمالي المصروف", "إجمالي السُلف", "إجمالي الخصومات",
        "إجمالي الراتب", "أشهر مصروفة", "أشهر مسجّلة", "المدرب", "م",
    ]
    data = [[Paragraph(h, hdr_style) for h in pdf_headers]]
    for i, r in enumerate(coach_rows, 1):
        data.append([
            Paragraph(f"{r['late_minutes_avg']}", cell_style),
            Paragraph(str(r["absent_days_total"]), cell_style),
            Paragraph(f"{r['pending_advances_total']:,.2f}", cell_style),
            Paragraph(f"{r['total_net_disbursed']:,.2f}", cell_style),
            Paragraph(f"{r['total_advances_repaid']:,.2f}", cell_style),
            Paragraph(f"{r['total_deductions']:,.2f}", cell_style),
            Paragraph(f"{r['total_base']:,.2f}", cell_style),
            Paragraph(str(r["months_disbursed"]), cell_style),
            Paragraph(str(r["months_recorded"]), cell_style),
            Paragraph(r["coach_name"], cell_style),
            Paragraph(str(i), cell_style),
        ])
    data.append([
        Paragraph("", cell_style),
        Paragraph(str(totals["absent_days_total"]), cell_style),
        Paragraph(f"{totals['pending_advances_total']:,.2f}", cell_style),
        Paragraph(f"{totals['total_net_disbursed']:,.2f}", cell_style),
        Paragraph(f"{totals['total_advances_repaid']:,.2f}", cell_style),
        Paragraph(f"{totals['total_deductions']:,.2f}", cell_style),
        Paragraph(f"{totals['total_base']:,.2f}", cell_style),
        Paragraph("", cell_style),
        Paragraph("", cell_style),
        Paragraph("الإجمالي", cell_style),
        Paragraph("", cell_style),
    ])

    col_widths_pdf = [22*mm, 18*mm, 24*mm, 26*mm, 24*mm, 26*mm, 26*mm, 18*mm, 18*mm, 42*mm, 8*mm]
    table = Table(data, colWidths=col_widths_pdf, repeatRows=1)
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#F97316')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('GRID', (0, 0), (-1, -1), 0.4, colors.HexColor('#DDDDDD')),
        ('ROWBACKGROUNDS', (0, 1), (-1, -2), [colors.white, colors.HexColor('#FFF7ED')]),
        ('BACKGROUND', (0, -1), (-1, -1), colors.HexColor('#FEF3C7')),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
    ]))
    elements.append(table)

    doc.build(elements)
    buffer.seek(0)
    filename_pdf = f"coach_salaries_report_{from_month}_{to_month}.pdf"
    return StreamingResponse(
        buffer,
        media_type="application/pdf",
        headers={"Content-Disposition": f"attachment; filename={filename_pdf}"}
    )


@router.get("/{salary_id}/payslip.pdf")
async def payslip_pdf(
    salary_id: str,
    current_user: dict = Depends(get_current_user)
):
    await require_permission(current_user, "salaries")
    effective_branch = resolve_branch_filter(current_user, None)
    query: dict = {"id": salary_id}
    if effective_branch:
        query.update(_branch_scope_filter(effective_branch))

    salary = await db.coach_salaries.find_one(query, {"_id": 0})
    if not salary:
        raise HTTPException(status_code=404, detail="سجل الراتب غير موجود")

    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    font_name = "Helvetica"
    for fp in [
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
        "/usr/share/fonts/TTF/DejaVuSans.ttf",
        "/nix/store/dejavu-fonts/share/fonts/truetype/DejaVuSans.ttf",
    ]:
        try:
            if Path(fp).exists():
                pdfmetrics.registerFont(TTFont('ArabicFont', fp))
                font_name = 'ArabicFont'
                break
        except Exception:
            continue

    title_style = ParagraphStyle('T', fontName=font_name, fontSize=14, leading=18, alignment=1)
    sub_style = ParagraphStyle('S', fontName=font_name, fontSize=10, leading=14, alignment=2)
    cell_style = ParagraphStyle('C', fontName=font_name, fontSize=10, leading=14, alignment=2)
    cell_center = ParagraphStyle('CC', fontName=font_name, fontSize=10, leading=14, alignment=1)
    hdr_style = ParagraphStyle('H', fontName=font_name, fontSize=10, leading=14,
                               textColor=colors.white, alignment=1)

    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4,
                            topMargin=15*mm, bottomMargin=15*mm,
                            leftMargin=15*mm, rightMargin=15*mm)
    elements = []
    elements.append(Paragraph("شركة اداء الابطال العالمية للرياضة", title_style))
    elements.append(Paragraph("Champions Academy", title_style))
    elements.append(Paragraph("قسيمة راتب — Payslip", title_style))
    elements.append(Spacer(1, 3*mm))
    elements.append(Paragraph(f"الشهر / Month: {salary['year_month']}", sub_style))
    elements.append(Paragraph(f"المدرب / Coach: {salary.get('coach_name', '')}", sub_style))
    ct = salary.get('contract_type', 'full_time')
    ct_label = 'دوام جزئي / Part-time' if ct == 'part_time' else 'دوام كامل / Full-time'
    elements.append(Paragraph(f"نوع التعاقد / Contract: {ct_label}", sub_style))
    elements.append(Paragraph(f"أيام العمل الشهرية / Monthly Work Days: {salary.get('monthly_work_days', 30)}", sub_style))
    status_label = 'مصروف / Disbursed' if salary.get('status') == 'disbursed' else 'مسودة / Draft'
    elements.append(Paragraph(f"الحالة / Status: {status_label}", sub_style))
    if salary.get("disbursed_at"):
        elements.append(Paragraph(f"تاريخ الصرف / Disbursed at: {salary['disbursed_at'][:10]}", sub_style))
    elements.append(Spacer(1, 5*mm))

    rows = [
        [Paragraph("Amount (SAR) / القيمة", hdr_style), Paragraph("Item / البند", hdr_style)],
        [Paragraph(f"{salary.get('base_salary', 0):,.2f}", cell_center), Paragraph("الراتب الأساسي / Base Salary", cell_style)],
        [Paragraph(f"{salary.get('monthly_work_days', 30)}", cell_center), Paragraph("أيام العمل المطلوبة / Required Work Days", cell_style)],
        [Paragraph(f"{salary.get('present_days', 0)}", cell_center), Paragraph("أيام الحضور / Present Days", cell_style)],
        [Paragraph(f"{salary.get('absent_days', 0)}", cell_center), Paragraph("أيام الغياب (تلقائي) / Absent Days (Auto)", cell_style)],
        [Paragraph(f"-{salary.get('deduction_absent', 0):,.2f}", cell_center), Paragraph("خصم الغياب / Absence Deduction", cell_style)],
        [Paragraph(f"{salary.get('late_minutes', 0)}", cell_center), Paragraph("دقائق التأخير / Late Minutes", cell_style)],
        [Paragraph(f"-{salary.get('deduction_late', 0):,.2f}", cell_center), Paragraph("خصم التأخير / Late Deduction", cell_style)],
        [Paragraph(f"+{salary.get('bonus', 0):,.2f}", cell_center), Paragraph("علاوة / Bonus", cell_style)],
    ]
    for m in (salary.get("manual_deductions") or []):
        rows.append([
            Paragraph(f"-{float(m.get('amount', 0)):,.2f}", cell_center),
            Paragraph(f"خصم / Deduction: {m.get('description', '')}", cell_style),
        ])
    rows.append([
        Paragraph(f"-{salary.get('advances_repaid_total', 0):,.2f}", cell_center),
        Paragraph("خصم سُلف / Advances Repaid", cell_style),
    ])
    rows.append([
        Paragraph(f"<b>{salary.get('net_amount', 0):,.2f}</b>", cell_center),
        Paragraph("<b>الصافي المستحق / Net Payable</b>", cell_style),
    ])

    t = Table(rows, colWidths=[60*mm, 110*mm])
    t.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#F97316')),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#DDDDDD')),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('ROWBACKGROUNDS', (0, 1), (-1, -2), [colors.white, colors.HexColor('#FFF7ED')]),
        ('BACKGROUND', (0, -1), (-1, -1), colors.HexColor('#FEF3C7')),
    ]))
    elements.append(t)
    elements.append(Spacer(1, 10*mm))

    sig = Table(
        [[Paragraph("توقيع المستلم / Recipient Signature", cell_center), Paragraph("توقيع المدير / Manager Signature", cell_center)],
         [Paragraph("____________________", cell_center), Paragraph("____________________", cell_center)]],
        colWidths=[85*mm, 85*mm]
    )
    sig.setStyle(TableStyle([('VALIGN', (0, 0), (-1, -1), 'MIDDLE')]))
    elements.append(sig)

    doc.build(elements)
    buffer.seek(0)
    fname = f"payslip_{salary['year_month']}_{salary.get('coach_name', '')}.pdf"
    return StreamingResponse(
        buffer,
        media_type="application/pdf",
        headers={"Content-Disposition": f"attachment; filename={fname}"}
    )
