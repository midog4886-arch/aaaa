"""
Coach Attendance API Routes
Handles coach/trainer attendance tracking (check-in, check-out, reports)
"""
from fastapi import APIRouter, HTTPException, Depends
from fastapi.responses import StreamingResponse
from pydantic import BaseModel
from typing import Optional, List
from datetime import datetime, timezone, timedelta
from io import BytesIO
from pathlib import Path
import uuid

from database import db
from utils.auth import get_current_user, require_branch_scope, resolve_branch_filter


def _branch_scope_filter(effective_branch: Optional[str]) -> dict:
    """Build a Mongo sub-filter that pins records to the caller's branch.

    Returns an empty dict for admins (no branch restriction). For non-admins,
    matches records whose ``branch_id`` equals the caller's branch OR which
    have no ``branch_id`` at all (legacy/unscoped records remain visible —
    same convention used by ``coaches`` and other branch-scoped routes).
    """
    if not effective_branch:
        return {}
    return {"$or": [
        {"branch_id": effective_branch},
        {"branch_id": None},
        {"branch_id": {"$exists": False}}
    ]}


def _get_openpyxl():
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
    return type('XL', (), {
        'Workbook': Workbook, 'Font': Font, 'Alignment': Alignment,
        'Border': Border, 'Side': Side, 'PatternFill': PatternFill,
    })()


def _get_reportlab():
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    return type('RL', (), {
        'colors': colors, 'A4': A4,
        'SimpleDocTemplate': SimpleDocTemplate, 'Table': Table,
        'TableStyle': TableStyle, 'Paragraph': Paragraph, 'Spacer': Spacer,
        'getSampleStyleSheet': getSampleStyleSheet, 'ParagraphStyle': ParagraphStyle,
        'mm': mm, 'pdfmetrics': pdfmetrics, 'TTFont': TTFont,
    })()

router = APIRouter(prefix="/coach-attendance", tags=["Coach Attendance"])

SAUDI_OFFSET = timedelta(hours=3)

def get_saudi_now():
    return datetime.now(timezone.utc) + SAUDI_OFFSET

class CheckInRequest(BaseModel):
    coach_id: str
    date: Optional[str] = None
    check_in_time: Optional[str] = None
    notes: Optional[str] = ""

class CheckOutRequest(BaseModel):
    check_out_time: Optional[str] = None

class MarkAbsentRequest(BaseModel):
    coach_id: str
    date: Optional[str] = None
    status: str = "absent"
    reason: Optional[str] = ""

class UpdateRecordRequest(BaseModel):
    check_in_time: Optional[str] = None
    check_out_time: Optional[str] = None
    notes: Optional[str] = None
    status: Optional[str] = None
    reason: Optional[str] = None


@router.get("")
async def get_coach_attendance(
    date: Optional[str] = None,
    coach_id: Optional[str] = None,
    month: Optional[str] = None,
    branch_filter: Optional[str] = None,
    current_user: dict = Depends(get_current_user)
):
    query = {}

    if date:
        query["date"] = date
    elif month:
        query["date"] = {"$regex": f"^{month}"}
    else:
        query["date"] = get_saudi_now().strftime("%Y-%m-%d")

    if coach_id:
        query["coach_id"] = coach_id

    # Branch filtering — fail-closed for non-admins without a branch_id.
    # Admins may pass branch_filter (None/"all" = no restriction); non-admins
    # are always pinned to their own branch regardless of what they send.
    effective_branch = resolve_branch_filter(current_user, branch_filter)
    if effective_branch:
        query.update(_branch_scope_filter(effective_branch))

    records = await db.coach_attendance.find(query, {"_id": 0}).sort("check_in_time", 1).to_list(500)
    return records


@router.post("/check-in")
async def check_in_coach(
    req: CheckInRequest,
    current_user: dict = Depends(get_current_user)
):
    now = get_saudi_now()
    date = req.date or now.strftime("%Y-%m-%d")
    check_in_time = req.check_in_time or now.strftime("%H:%M")

    # Fail-closed for non-admins without a branch_id (HTTP 403).
    effective_branch = require_branch_scope(current_user)

    coach = await db.coaches.find_one({"id": req.coach_id}, {"_id": 0})
    if not coach:
        raise HTTPException(status_code=404, detail="Coach not found")

    # Prevent IDOR: non-admins may only check in coaches in their own branch.
    # Coaches with no branch_id (legacy) remain visible to any branch.
    if effective_branch:
        coach_branch = coach.get("branch_id")
        if coach_branch and coach_branch != effective_branch:
            raise HTTPException(status_code=404, detail="Coach not found")

    existing = await db.coach_attendance.find_one({
        "coach_id": req.coach_id,
        "date": date,
        "status": {"$in": ["present", "checked_out"]}
    })
    if existing:
        raise HTTPException(status_code=400, detail="Coach already checked in for this date")

    await db.coach_attendance.delete_many({
        "coach_id": req.coach_id,
        "date": date,
        "status": {"$in": ["absent", "leave"]}
    })

    record = {
        "id": str(uuid.uuid4()),
        "coach_id": req.coach_id,
        "coach_name": coach.get("name_ar", coach.get("name", "")),
        "coach_phone": coach.get("phone", ""),
        "date": date,
        "check_in_time": check_in_time,
        "check_out_time": None,
        "total_hours": None,
        "status": "present",
        "notes": req.notes or "",
        "reason": "",
        "branch_id": coach.get("branch_id"),
        "recorded_by": current_user.get("name", current_user.get("username", "")),
        "created_at": now.isoformat()
    }

    await db.coach_attendance.insert_one(record)
    record.pop("_id", None)
    return record


@router.post("/{record_id}/check-out")
async def check_out_coach(
    record_id: str,
    req: CheckOutRequest,
    current_user: dict = Depends(get_current_user)
):
    now = get_saudi_now()
    check_out_time = req.check_out_time or now.strftime("%H:%M")

    record = await db.coach_attendance.find_one({"id": record_id}, {"_id": 0})
    if not record:
        raise HTTPException(status_code=404, detail="Attendance record not found")

    if record.get("status") == "checked_out":
        raise HTTPException(status_code=400, detail="Coach already checked out")

    total_hours = None
    if record.get("check_in_time"):
        try:
            cin = datetime.strptime(record["check_in_time"], "%H:%M")
            cout = datetime.strptime(check_out_time, "%H:%M")
            diff = (cout - cin).total_seconds() / 3600
            if diff < 0:
                diff += 24
            total_hours = round(diff, 2)
        except:
            pass

    await db.coach_attendance.update_one(
        {"id": record_id},
        {"$set": {
            "check_out_time": check_out_time,
            "total_hours": total_hours,
            "status": "checked_out"
        }}
    )

    return {
        "message": "Check-out recorded",
        "check_out_time": check_out_time,
        "total_hours": total_hours
    }


@router.post("/mark-absent")
async def mark_absent(
    req: MarkAbsentRequest,
    current_user: dict = Depends(get_current_user)
):
    now = get_saudi_now()
    date = req.date or now.strftime("%Y-%m-%d")

    # Fail-closed for non-admins without a branch_id (HTTP 403).
    effective_branch = require_branch_scope(current_user)

    coach = await db.coaches.find_one({"id": req.coach_id}, {"_id": 0})
    if not coach:
        raise HTTPException(status_code=404, detail="Coach not found")

    # Prevent IDOR: non-admins may only mark coaches in their own branch.
    # Coaches with no branch_id (legacy) remain visible to any branch.
    if effective_branch:
        coach_branch = coach.get("branch_id")
        if coach_branch and coach_branch != effective_branch:
            raise HTTPException(status_code=404, detail="Coach not found")

    existing = await db.coach_attendance.find_one({
        "coach_id": req.coach_id,
        "date": date,
        "status": {"$in": ["present", "checked_out"]}
    })
    if existing:
        raise HTTPException(status_code=400, detail="Coach already has attendance for this date")

    await db.coach_attendance.delete_many({
        "coach_id": req.coach_id,
        "date": date,
        "status": {"$in": ["absent", "leave"]}
    })

    record = {
        "id": str(uuid.uuid4()),
        "coach_id": req.coach_id,
        "coach_name": coach.get("name_ar", coach.get("name", "")),
        "coach_phone": coach.get("phone", ""),
        "date": date,
        "check_in_time": None,
        "check_out_time": None,
        "total_hours": None,
        "status": req.status,
        "notes": "",
        "reason": req.reason or "",
        "branch_id": coach.get("branch_id"),
        "recorded_by": current_user.get("name", current_user.get("username", "")),
        "created_at": now.isoformat()
    }

    await db.coach_attendance.insert_one(record)
    record.pop("_id", None)
    return record


@router.put("/{record_id}")
async def update_record(
    record_id: str,
    req: UpdateRecordRequest,
    current_user: dict = Depends(get_current_user)
):
    # Fail-closed for non-admins without a branch_id (HTTP 403).
    effective_branch = require_branch_scope(current_user)

    record = await db.coach_attendance.find_one({"id": record_id})
    if not record:
        raise HTTPException(status_code=404, detail="Record not found")

    # Prevent IDOR: non-admins may only edit records in their own branch.
    # Legacy records with no branch_id remain editable by any branch.
    if effective_branch:
        rec_branch = record.get("branch_id")
        if rec_branch and rec_branch != effective_branch:
            raise HTTPException(status_code=404, detail="Record not found")

    updates = {}
    if req.check_in_time is not None:
        updates["check_in_time"] = req.check_in_time
    if req.check_out_time is not None:
        updates["check_out_time"] = req.check_out_time
    if req.notes is not None:
        updates["notes"] = req.notes
    if req.status is not None:
        updates["status"] = req.status
    if req.reason is not None:
        updates["reason"] = req.reason

    cin = req.check_in_time or record.get("check_in_time")
    cout = req.check_out_time or record.get("check_out_time")
    if cin and cout:
        try:
            c1 = datetime.strptime(cin, "%H:%M")
            c2 = datetime.strptime(cout, "%H:%M")
            diff = (c2 - c1).total_seconds() / 3600
            if diff < 0:
                diff += 24
            updates["total_hours"] = round(diff, 2)
        except:
            pass

    if updates:
        await db.coach_attendance.update_one({"id": record_id}, {"$set": updates})

    updated = await db.coach_attendance.find_one({"id": record_id}, {"_id": 0})
    try:
        from utils.audit import log_audit
        await log_audit(
            actor=current_user,
            action="coach_attendance.update",
            entity_type="coach_attendance",
            entity_id=record_id,
            entity_name=record.get("coach_name", record.get("coach_id", "")),
            before=record,
            after=updated,
        )
    except Exception:
        pass
    return updated


@router.delete("/{record_id}")
async def delete_record(
    record_id: str,
    current_user: dict = Depends(get_current_user)
):
    # Fail-closed for non-admins without a branch_id (HTTP 403).
    effective_branch = require_branch_scope(current_user)

    # Prevent IDOR: scope the delete query so non-admins cannot remove
    # records from other branches. Legacy records with no branch_id remain
    # deletable by any branch (matches read scoping convention).
    delete_query: dict = {"id": record_id}
    if effective_branch:
        delete_query["$or"] = [
            {"branch_id": effective_branch},
            {"branch_id": None},
            {"branch_id": {"$exists": False}},
        ]

    before = await db.coach_attendance.find_one(delete_query, {"_id": 0})
    result = await db.coach_attendance.delete_one(delete_query)
    if result.deleted_count == 0:
        raise HTTPException(status_code=404, detail="Record not found")
    try:
        from utils.audit import log_audit
        await log_audit(
            actor=current_user,
            action="coach_attendance.delete",
            entity_type="coach_attendance",
            entity_id=record_id,
            entity_name=(before or {}).get("coach_name") or (before or {}).get("coach_id", ""),
            before=before,
        )
    except Exception:
        pass
    return {"message": "Record deleted"}


@router.get("/monthly-report")
async def monthly_report(
    month: str,
    branch_filter: Optional[str] = None,
    late_threshold: str = "09:00",
    current_user: dict = Depends(get_current_user)
):
    # Branch filtering — fail-closed for non-admins without a branch_id.
    # Admins may pass branch_filter (None/"all" = no restriction); non-admins
    # are always pinned to their own branch regardless of what they send.
    effective_branch = resolve_branch_filter(current_user, branch_filter)

    query: dict = {"date": {"$regex": f"^{month}"}}
    if effective_branch:
        query.update(_branch_scope_filter(effective_branch))

    records = await db.coach_attendance.find(query, {"_id": 0}).to_list(5000)

    coach_query = _branch_scope_filter(effective_branch) if effective_branch else {}
    coaches = await db.coaches.find(coach_query, {"_id": 0}).to_list(100)

    # Parse late threshold once; fall back to 09:00 on invalid input
    try:
        threshold_dt = datetime.strptime(late_threshold, "%H:%M")
        effective_threshold = late_threshold
    except Exception:
        threshold_dt = datetime.strptime("09:00", "%H:%M")
        effective_threshold = "09:00"

    report = {}
    for coach in coaches:
        cid = coach["id"]
        coach_records = [r for r in records if r.get("coach_id") == cid]
        present_days = len([r for r in coach_records if r.get("status") in ("present", "checked_out")])
        absent_days = len([r for r in coach_records if r.get("status") == "absent"])
        leave_days = len([r for r in coach_records if r.get("status") == "leave"])
        total_hours = sum(r.get("total_hours", 0) or 0 for r in coach_records)

        # ── Late arrivals (per-coach threshold or global fallback) ──
        # Use the coach's own expected_checkin_time if set, else the global threshold
        coach_threshold_str = coach.get("expected_checkin_time") or effective_threshold
        try:
            coach_threshold_dt = datetime.strptime(coach_threshold_str, "%H:%M")
        except Exception:
            coach_threshold_dt = threshold_dt

        late_records = []
        late_days = 0
        late_minutes_total = 0
        for r in coach_records:
            if r.get("status") not in ("present", "checked_out"):
                continue
            cin_str = r.get("check_in_time")
            if not cin_str:
                continue
            try:
                cin_dt = datetime.strptime(cin_str, "%H:%M")
                diff = (cin_dt - coach_threshold_dt).total_seconds() / 60
                if diff > 0:
                    late_days += 1
                    late_minutes_total += diff
                    late_records.append({
                        "date": r.get("date", ""),
                        "check_in_time": cin_str,
                        "minutes_late": round(diff),
                        "threshold": coach_threshold_str
                    })
            except Exception:
                pass
        # Sort late records by date
        late_records.sort(key=lambda x: x["date"])

        contract_type = coach.get("contract_type") or "full_time"
        if contract_type not in ("full_time", "part_time"):
            contract_type = "full_time"
        try:
            monthly_work_days = max(1, min(int(coach.get("monthly_work_days") or 30), 31))
        except (ValueError, TypeError):
            monthly_work_days = 30

        attendance_percentage = round((present_days / monthly_work_days) * 100) if monthly_work_days > 0 else 0

        report[cid] = {
            "coach_id": cid,
            "coach_name": coach.get("name_ar", coach.get("name", "")),
            "contract_type": contract_type,
            "monthly_work_days": monthly_work_days,
            "attendance_percentage": attendance_percentage,
            "present_days": present_days,
            "absent_days": absent_days,
            "leave_days": leave_days,
            "total_hours": round(total_hours, 2),
            "late_days": late_days,
            "late_minutes": round(late_minutes_total),
            "late_records": late_records,
            "coach_threshold": coach_threshold_str,
            "records": coach_records
        }

    return {"month": month, "late_threshold": effective_threshold, "report": list(report.values())}


@router.get("/monthly-report/export")
async def export_monthly_report(
    month: str,
    branch_filter: Optional[str] = None,
    late_threshold: str = "09:00",
    format: str = "xlsx",
    current_user: dict = Depends(get_current_user)
):
    """Export monthly coach attendance report as Excel or PDF."""
    if not current_user.get("is_admin"):
        raise HTTPException(status_code=403, detail="غير مصرح: التصدير متاح للمدير فقط")
    if format not in ("xlsx", "pdf"):
        raise HTTPException(status_code=400, detail="format يجب أن يكون xlsx أو pdf")

    # Branch filtering — fail-closed for non-admins without a branch_id.
    # Admins may pass branch_filter (None/"all" = no restriction); non-admins
    # are always pinned to their own branch regardless of what they send.
    effective_branch = resolve_branch_filter(current_user, branch_filter)

    query: dict = {"date": {"$regex": f"^{month}"}}
    if effective_branch:
        query.update(_branch_scope_filter(effective_branch))
    records = await db.coach_attendance.find(query, {"_id": 0}).to_list(5000)

    coach_query = _branch_scope_filter(effective_branch) if effective_branch else {}
    coaches = await db.coaches.find(coach_query, {"_id": 0}).to_list(100)

    try:
        threshold_dt = datetime.strptime(late_threshold, "%H:%M")
        effective_threshold = late_threshold
    except Exception:
        threshold_dt = datetime.strptime("09:00", "%H:%M")
        effective_threshold = "09:00"

    rows = []
    for idx, coach in enumerate(coaches, 1):
        cid = coach["id"]
        coach_records = [r for r in records if r.get("coach_id") == cid]
        present_days = len([r for r in coach_records if r.get("status") in ("present", "checked_out")])
        absent_days = len([r for r in coach_records if r.get("status") == "absent"])
        leave_days = len([r for r in coach_records if r.get("status") == "leave"])
        total_hours = round(sum(r.get("total_hours", 0) or 0 for r in coach_records), 2)

        coach_threshold_str = coach.get("expected_checkin_time") or effective_threshold
        try:
            coach_threshold_dt = datetime.strptime(coach_threshold_str, "%H:%M")
        except Exception:
            coach_threshold_dt = threshold_dt

        late_days = 0
        late_minutes_total = 0
        for r in coach_records:
            if r.get("status") not in ("present", "checked_out"):
                continue
            cin_str = r.get("check_in_time")
            if not cin_str:
                continue
            try:
                cin_dt = datetime.strptime(cin_str, "%H:%M")
                diff = (cin_dt - coach_threshold_dt).total_seconds() / 60
                if diff > 0:
                    late_days += 1
                    late_minutes_total += diff
            except Exception:
                pass

        contract_type = coach.get("contract_type") or "full_time"
        if contract_type not in ("full_time", "part_time"):
            contract_type = "full_time"
        try:
            monthly_work_days = max(1, min(int(coach.get("monthly_work_days") or 30), 31))
        except (ValueError, TypeError):
            monthly_work_days = 30

        attendance_percentage = round((present_days / monthly_work_days) * 100) if monthly_work_days > 0 else 0

        rows.append({
            "idx": idx,
            "coach_name": coach.get("name_ar", coach.get("name", "")),
            "contract_type": contract_type,
            "monthly_work_days": monthly_work_days,
            "attendance_percentage": attendance_percentage,
            "present_days": present_days,
            "absent_days": absent_days,
            "leave_days": leave_days,
            "total_hours": total_hours,
            "late_days": late_days,
            "late_minutes": round(late_minutes_total),
        })

    export_date = datetime.now().strftime("%Y-%m-%d")
    headers_row = ["م", "المدرب", "نوع التعاقد", "أيام العمل", "أيام الحضور", "نسبة الحضور", "أيام الغياب", "أيام الإجازة", "إجمالي الساعات", "أيام التأخر", "دقائق التأخر"]

    # ── Excel ──────────────────────────────────────────
    if format == "xlsx":
        xl = _get_openpyxl()
        Workbook = xl.Workbook; Font = xl.Font; PatternFill = xl.PatternFill
        Border = xl.Border; Side = xl.Side; Alignment = xl.Alignment
        wb = Workbook()
        ws = wb.active
        ws.title = "تقرير المدربين"

        orange_fill = PatternFill("solid", fgColor="F97316")
        white_on_orange = Font(bold=True, color="FFFFFF")
        alt_fill = PatternFill("solid", fgColor="FFF7ED")
        border = Border(
            left=Side(style='thin', color='DDDDDD'),
            right=Side(style='thin', color='DDDDDD'),
            top=Side(style='thin', color='DDDDDD'),
            bottom=Side(style='thin', color='DDDDDD'),
        )
        center = Alignment(horizontal='center', vertical='center', wrap_text=True)
        right_align = Alignment(horizontal='right', vertical='center', wrap_text=True)

        ws.append(["شركة اداء الابطال العالمية للرياضة"])
        ws.append([f"التقرير الشهري للمدربين — {month}"])
        ws.append([f"تاريخ التصدير: {export_date}"])
        ws.append([])

        col_count = len(headers_row)
        last_col_letter = chr(ord('A') + col_count - 1)
        for col in range(1, col_count + 1):
            ws.cell(row=1, column=col).font = Font(bold=True, size=13)
        ws.merge_cells(f'A1:{last_col_letter}1')
        ws.merge_cells(f'A2:{last_col_letter}2')
        ws.merge_cells(f'A3:{last_col_letter}3')
        ws['A1'].alignment = right_align
        ws['A2'].alignment = right_align
        ws['A3'].alignment = right_align

        ws.append(headers_row)
        header_row_num = 5
        for col_idx, _ in enumerate(headers_row, 1):
            cell = ws.cell(row=header_row_num, column=col_idx)
            cell.fill = orange_fill
            cell.font = white_on_orange
            cell.alignment = center
            cell.border = border

        contract_labels = {"full_time": "دوام كامل", "part_time": "دوام جزئي"}
        green_pct_font = Font(bold=True, color="166534")
        yellow_pct_font = Font(bold=True, color="92400E")
        red_pct_font = Font(bold=True, color="991B1B")
        for data_idx, r in enumerate(rows):
            pct = r["attendance_percentage"]
            ws.append([r["idx"], r["coach_name"],
                       contract_labels.get(r["contract_type"], r["contract_type"]),
                       r["monthly_work_days"],
                       r["present_days"], f"{pct}%", r["absent_days"],
                       r["leave_days"], r["total_hours"], r["late_days"], r["late_minutes"]])
            row_num = header_row_num + 1 + data_idx
            use_alt = data_idx % 2 == 1
            for col_idx in range(1, col_count + 1):
                cell = ws.cell(row=row_num, column=col_idx)
                if use_alt:
                    cell.fill = alt_fill
                cell.border = border
                cell.alignment = center if col_idx != 2 else right_align
            pct_cell = ws.cell(row=row_num, column=6)
            if pct >= 80:
                pct_cell.font = green_pct_font
            elif pct >= 50:
                pct_cell.font = yellow_pct_font
            else:
                pct_cell.font = red_pct_font

        col_widths = [6, 28, 14, 12, 14, 14, 14, 14, 16, 14, 16]
        for idx_w, w in enumerate(col_widths, 1):
            ws.column_dimensions[ws.cell(row=1, column=idx_w).column_letter].width = w

        buffer = BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        filename = f"coach_report_{month}.xlsx"
        return StreamingResponse(
            buffer,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f"attachment; filename={filename}"}
        )

    # ── PDF ────────────────────────────────────────────
    rl = _get_reportlab()
    colors = rl.colors; A4 = rl.A4; mm = rl.mm
    SimpleDocTemplate = rl.SimpleDocTemplate; Table = rl.Table
    TableStyle = rl.TableStyle; Paragraph = rl.Paragraph; Spacer = rl.Spacer
    getSampleStyleSheet = rl.getSampleStyleSheet
    ParagraphStyle = rl.ParagraphStyle
    pdfmetrics = rl.pdfmetrics; TTFont = rl.TTFont

    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4,
                            topMargin=15*mm, bottomMargin=15*mm,
                            leftMargin=15*mm, rightMargin=15*mm)

    font_name = "Helvetica"
    font_paths = [
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
        "/usr/share/fonts/TTF/DejaVuSans.ttf",
        "/nix/store/dejavu-fonts/share/fonts/truetype/DejaVuSans.ttf",
    ]
    for fp in font_paths:
        try:
            if Path(fp).exists():
                pdfmetrics.registerFont(TTFont('ArabicFont', fp))
                font_name = 'ArabicFont'
                break
        except Exception:
            continue

    title_style = ParagraphStyle('T', fontName=font_name, fontSize=13, leading=18, alignment=2)
    sub_style = ParagraphStyle('S', fontName=font_name, fontSize=9, leading=13,
                               textColor=colors.HexColor('#555555'), alignment=2)
    cell_style = ParagraphStyle('C', fontName=font_name, fontSize=8, leading=11, alignment=1)
    hdr_style = ParagraphStyle('H', fontName=font_name, fontSize=8, leading=11,
                               textColor=colors.white, alignment=1)

    elements = []
    elements.append(Paragraph("شركة اداء الابطال العالمية للرياضة", title_style))
    elements.append(Paragraph(f"التقرير الشهري للمدربين — {month}", title_style))
    elements.append(Paragraph(f"تاريخ التصدير: {export_date}", sub_style))
    elements.append(Spacer(1, 5*mm))

    contract_labels = {"full_time": "دوام كامل", "part_time": "دوام جزئي"}
    pdf_headers = ["دقائق التأخر", "أيام التأخر", "الساعات", "الإجازة", "الغياب", "نسبة الحضور", "الحضور", "أيام العمل", "التعاقد", "المدرب", "م"]
    header_row_pdf = [Paragraph(h, hdr_style) for h in pdf_headers]
    data = [header_row_pdf]
    for row_idx, r in enumerate(rows, start=1):
        pct = r["attendance_percentage"]
        if pct >= 80:
            pct_hex = '#166534'
        elif pct >= 50:
            pct_hex = '#92400E'
        else:
            pct_hex = '#991B1B'
        pct_style = ParagraphStyle(f'P{row_idx}', fontName=font_name, fontSize=8, leading=11,
                                    alignment=1, textColor=colors.HexColor(pct_hex))
        data.append([
            Paragraph(str(r["late_minutes"]), cell_style),
            Paragraph(str(r["late_days"]), cell_style),
            Paragraph(str(r["total_hours"]), cell_style),
            Paragraph(str(r["leave_days"]), cell_style),
            Paragraph(str(r["absent_days"]), cell_style),
            Paragraph(f"<b>{pct}%</b>", pct_style),
            Paragraph(str(r["present_days"]), cell_style),
            Paragraph(str(r["monthly_work_days"]), cell_style),
            Paragraph(contract_labels.get(r["contract_type"], r["contract_type"]), cell_style),
            Paragraph(r["coach_name"], cell_style),
            Paragraph(str(r["idx"]), cell_style),
        ])

    col_widths_pdf = [16*mm, 14*mm, 14*mm, 14*mm, 14*mm, 16*mm, 14*mm, 14*mm, 16*mm, 32*mm, 8*mm]
    table = Table(data, colWidths=col_widths_pdf, repeatRows=1)
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#F97316')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('GRID', (0, 0), (-1, -1), 0.4, colors.HexColor('#DDDDDD')),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#FFF7ED')]),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
    ]))
    elements.append(table)
    elements.append(Spacer(1, 5*mm))
    elements.append(Paragraph(f"إجمالي المدربين: {len(rows)}", sub_style))

    doc.build(elements)
    buffer.seek(0)
    filename_pdf = f"coach_report_{month}.pdf"
    return StreamingResponse(
        buffer,
        media_type="application/pdf",
        headers={"Content-Disposition": f"attachment; filename={filename_pdf}"}
    )


# ══════════════════════════════════════════════════════
#  PUBLIC QR CHECK-IN / CHECK-OUT  (no auth required)
# ══════════════════════════════════════════════════════

# ── Helper: get coach by employee_id ─────────────────
async def get_coach_by_employee_id(employee_id: str):
    return await db.coaches.find_one({"employee_id": employee_id}, {"_id": 0})


@router.get("/qr-status-by-code/{employee_id}")
async def get_coach_qr_status_by_code(employee_id: str):
    """Public endpoint: look up coach by employee_id (numeric code) and return today's status."""
    coach = await get_coach_by_employee_id(employee_id)
    if not coach:
        raise HTTPException(status_code=404, detail="المدرب غير موجود")

    coach_id = coach["id"]
    today = get_saudi_now().strftime("%Y-%m-%d")
    record = await db.coach_attendance.find_one(
        {"coach_id": coach_id, "date": today},
        {"_id": 0}
    )

    return {
        "coach_id": coach_id,
        "employee_id": employee_id,
        "coach_name": coach.get("name_ar", coach.get("name", "")),
        "coach_phone": coach.get("phone", ""),
        "today": today,
        "status": record.get("status") if record else None,
        "check_in_time": record.get("check_in_time") if record else None,
        "check_out_time": record.get("check_out_time") if record else None,
        "total_hours": record.get("total_hours") if record else None,
        "record_id": record.get("id") if record else None,
        "is_coach": True,
    }


@router.post("/qr-checkin-by-code/{employee_id}")
async def qr_checkin_coach_by_code(employee_id: str):
    """Public endpoint: check in/out coach by employee_id (numeric QR code)."""
    coach = await get_coach_by_employee_id(employee_id)
    if not coach:
        raise HTTPException(status_code=404, detail="المدرب غير موجود")

    coach_id = coach["id"]
    now = get_saudi_now()
    today = now.strftime("%Y-%m-%d")
    current_time = now.strftime("%H:%M")
    coach_name = coach.get("name_ar", coach.get("name", ""))

    record = await db.coach_attendance.find_one({"coach_id": coach_id, "date": today})

    if record and record.get("status") == "checked_out":
        return {
            "action": "already_out",
            "coach_name": coach_name,
            "employee_id": employee_id,
            "check_in_time": record.get("check_in_time"),
            "check_out_time": record.get("check_out_time"),
            "total_hours": record.get("total_hours"),
            "message": f"تم تسجيل انصرافك مسبقاً في {record.get('check_out_time', '')}",
            "is_coach": True,
        }

    if record and record.get("status") == "present":
        total_hours = None
        try:
            cin = datetime.strptime(record["check_in_time"], "%H:%M")
            cout = datetime.strptime(current_time, "%H:%M")
            diff = (cout - cin).total_seconds() / 3600
            if diff < 0:
                diff += 24
            total_hours = round(diff, 2)
        except Exception:
            pass

        await db.coach_attendance.update_one(
            {"id": record["id"]},
            {"$set": {"check_out_time": current_time, "total_hours": total_hours, "status": "checked_out"}}
        )
        return {
            "action": "checked_out",
            "coach_name": coach_name,
            "employee_id": employee_id,
            "check_in_time": record.get("check_in_time"),
            "check_out_time": current_time,
            "total_hours": total_hours,
            "message": f"تم تسجيل انصرافك بنجاح — {current_time}",
            "is_coach": True,
        }

    # Check in
    await db.coach_attendance.delete_many({
        "coach_id": coach_id, "date": today,
        "status": {"$in": ["absent", "leave"]}
    })
    new_record = {
        "id": str(uuid.uuid4()),
        "coach_id": coach_id,
        "coach_name": coach_name,
        "coach_phone": coach.get("phone", ""),
        "date": today,
        "check_in_time": current_time,
        "check_out_time": None,
        "total_hours": None,
        "status": "present",
        "notes": "تسجيل عبر QR",
        "reason": "",
        "branch_id": coach.get("branch_id"),
        "recorded_by": "QR",
        "created_at": now.isoformat()
    }
    await db.coach_attendance.insert_one(new_record)
    return {
        "action": "checked_in",
        "coach_name": coach_name,
        "employee_id": employee_id,
        "check_in_time": current_time,
        "check_out_time": None,
        "total_hours": None,
        "message": f"تم تسجيل حضورك بنجاح — {current_time}",
        "is_coach": True,
    }


@router.get("/qr-status/{coach_id}")
async def get_coach_qr_status(coach_id: str):
    """Public endpoint: returns coach name + today's attendance status for QR scan page."""
    coach = await db.coaches.find_one({"id": coach_id}, {"_id": 0})
    if not coach:
        raise HTTPException(status_code=404, detail="المدرب غير موجود")

    today = get_saudi_now().strftime("%Y-%m-%d")
    record = await db.coach_attendance.find_one(
        {"coach_id": coach_id, "date": today},
        {"_id": 0}
    )

    return {
        "coach_id": coach_id,
        "coach_name": coach.get("name_ar", coach.get("name", "")),
        "coach_phone": coach.get("phone", ""),
        "today": today,
        "status": record.get("status") if record else None,
        "check_in_time": record.get("check_in_time") if record else None,
        "check_out_time": record.get("check_out_time") if record else None,
        "total_hours": record.get("total_hours") if record else None,
        "record_id": record.get("id") if record else None,
    }


@router.post("/qr-checkin/{coach_id}")
async def qr_checkin_coach(coach_id: str):
    """Public endpoint: check in (if not yet) or check out (if present) via QR scan."""
    now = get_saudi_now()
    today = now.strftime("%Y-%m-%d")
    current_time = now.strftime("%H:%M")

    coach = await db.coaches.find_one({"id": coach_id}, {"_id": 0})
    if not coach:
        raise HTTPException(status_code=404, detail="المدرب غير موجود")

    coach_name = coach.get("name_ar", coach.get("name", ""))

    # Look for today's record
    record = await db.coach_attendance.find_one({"coach_id": coach_id, "date": today})

    # ── Already checked out ──────────────────────────────
    if record and record.get("status") == "checked_out":
        return {
            "action": "already_out",
            "coach_name": coach_name,
            "check_in_time": record.get("check_in_time"),
            "check_out_time": record.get("check_out_time"),
            "total_hours": record.get("total_hours"),
            "message": f"تم تسجيل انصرافك مسبقاً في {record.get('check_out_time', '')}"
        }

    # ── Present → Check out ──────────────────────────────
    if record and record.get("status") == "present":
        total_hours = None
        try:
            cin = datetime.strptime(record["check_in_time"], "%H:%M")
            cout = datetime.strptime(current_time, "%H:%M")
            diff = (cout - cin).total_seconds() / 3600
            if diff < 0:
                diff += 24
            total_hours = round(diff, 2)
        except Exception:
            pass

        await db.coach_attendance.update_one(
            {"id": record["id"]},
            {"$set": {
                "check_out_time": current_time,
                "total_hours": total_hours,
                "status": "checked_out"
            }}
        )
        return {
            "action": "checked_out",
            "coach_name": coach_name,
            "check_in_time": record.get("check_in_time"),
            "check_out_time": current_time,
            "total_hours": total_hours,
            "message": f"تم تسجيل انصرافك بنجاح — {current_time}"
        }

    # ── Not present → Check in ───────────────────────────
    # Remove any absent/leave records for today first
    await db.coach_attendance.delete_many({
        "coach_id": coach_id,
        "date": today,
        "status": {"$in": ["absent", "leave"]}
    })

    new_record = {
        "id": str(uuid.uuid4()),
        "coach_id": coach_id,
        "coach_name": coach_name,
        "coach_phone": coach.get("phone", ""),
        "date": today,
        "check_in_time": current_time,
        "check_out_time": None,
        "total_hours": None,
        "status": "present",
        "notes": "تسجيل عبر QR",
        "reason": "",
        "branch_id": coach.get("branch_id"),
        "recorded_by": "QR",
        "created_at": now.isoformat()
    }
    await db.coach_attendance.insert_one(new_record)
    new_record.pop("_id", None)

    return {
        "action": "checked_in",
        "coach_name": coach_name,
        "check_in_time": current_time,
        "check_out_time": None,
        "total_hours": None,
        "message": f"تم تسجيل حضورك بنجاح — {current_time}"
    }
