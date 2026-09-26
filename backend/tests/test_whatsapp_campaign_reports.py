import asyncio
from datetime import datetime, timezone

import pytest
from openpyxl import load_workbook

from routes import whatsapp as whatsapp_mod
from services import whatsapp_bulk_jobs as jobs


def run(coro):
    return asyncio.run(coro)


def _matches(row, query):
    for key, expected in query.items():
        if isinstance(expected, dict):
            if "$exists" in expected and ((key in row) != expected["$exists"]):
                return False
            if "$ne" in expected and row.get(key) == expected["$ne"]:
                return False
        elif row.get(key) != expected:
            return False
    return True


class Cursor:
    def __init__(self, rows):
        self.rows = [dict(row) for row in rows]

    async def to_list(self, length=None):
        return self.rows[:length] if length else self.rows

    def __aiter__(self):
        return self

    async def __anext__(self):
        if not self.rows:
            raise StopAsyncIteration
        return self.rows.pop(0)


class Collection:
    def __init__(self, rows=None):
        self.rows = list(rows or [])

    async def find_one(self, query, projection=None):
        row = next((row for row in self.rows if _matches(row, query)), None)
        return dict(row) if row else None

    def find(self, query, projection=None):
        return Cursor(row for row in self.rows if _matches(row, query))

    async def update_many(self, query, update):
        count = 0
        for row in self.rows:
            if _matches(row, query):
                row.update(update.get("$set", {}))
                count += 1
        return type("Result", (), {"matched_count": count})()


class DB:
    def __init__(self, jobs_rows=None, item_rows=None):
        self.data = {
            "whatsapp_campaign_jobs": Collection(jobs_rows),
            "whatsapp_campaign_job_items": Collection(item_rows),
            "members": Collection(),
            "users": Collection(),
            "branches": Collection(),
        }

    def __getitem__(self, name):
        return self.data.setdefault(name, Collection())


class ReceiptCollection(Collection):
    """Small Mongo-like store for receipt-buffer race regression coverage."""
    async def create_index(self, *args, **kwargs):
        return "index"

    async def update_one(self, query, update, upsert=False):
        row = next((row for row in self.rows if _matches(row, query)), None)
        if row:
            row.update(update.get("$set", {}))
            return type("Result", (), {
                "matched_count": 1, "modified_count": 1, "upserted_id": None,
            })()
        if upsert:
            row = {
                key: value for key, value in query.items()
                if not isinstance(value, dict)
            }
            row.update(update.get("$setOnInsert", {}))
            row.update(update.get("$set", {}))
            self.rows.append(row)
            return type("Result", (), {
                "matched_count": 0, "modified_count": 0, "upserted_id": "new",
            })()
        return type("Result", (), {
            "matched_count": 0, "modified_count": 0, "upserted_id": None,
        })()

    async def delete_one(self, query):
        for index, row in enumerate(self.rows):
            if _matches(row, query):
                self.rows.pop(index)
                return type("Result", (), {"deleted_count": 1})()
        return type("Result", (), {"deleted_count": 0})()


class ReceiptDB(DB):
    def __init__(self, item_rows=None):
        super().__init__(item_rows=item_rows)
        self.data["whatsapp_campaign_job_items"] = ReceiptCollection(item_rows)
        self.data[jobs.RECEIPT_BUFFER_COLLECTION] = ReceiptCollection()


@pytest.mark.parametrize("provider", ["meta_cloud", "waha", "whatsflow"])
def test_receipt_before_send_response_reconciles_exact_provider_id(monkeypatch, provider):
    db = ReceiptDB()
    monkeypatch.setattr(jobs, "_db", db)
    received_at = "2026-01-01T01:00:00+00:00"

    # The authenticated callback wins the race with the send response. There
    # is no item mapping yet, so only durable provider-ID evidence is kept.
    assert run(jobs.record_receipt(
        "branch-a", provider, "provider-id", "delivered",
        timestamp=received_at, tenant_slug="tenant-a",
    )) == 0
    assert db[jobs.RECEIPT_BUFFER_COLLECTION].rows[0]["provider_message_id"] == "provider-id"

    # Simulate the send response saving its exact ID after the callback.
    db["whatsapp_campaign_job_items"].rows.append({
        "id": "item-1", "branch_id": "branch-a", "provider": provider,
        "provider_message_id": "provider-id", "status": "sent",
    })
    assert run(jobs.reconcile_receipt(
        "branch-a", provider, "provider-id", tenant_slug="tenant-a",
    )) == 1
    item = db["whatsapp_campaign_job_items"].rows[0]
    assert item["delivery_status"] == "delivered"
    assert item["delivered_at"] == received_at
    assert not db[jobs.RECEIPT_BUFFER_COLLECTION].rows


def test_response_saved_between_receipt_lookup_and_buffer_write_is_not_lost(monkeypatch):
    db = ReceiptDB()
    monkeypatch.setattr(jobs, "_db", db)
    buffer = db[jobs.RECEIPT_BUFFER_COLLECTION]
    original_update = buffer.update_one
    saved = False

    async def write_receipt_then_save_response(query, update, upsert=False):
        nonlocal saved
        result = await original_update(query, update, upsert)
        if upsert and not saved:
            saved = True
            db["whatsapp_campaign_job_items"].rows.append({
                "id": "item-1", "branch_id": "branch-a",
                "provider": "meta_cloud", "provider_message_id": "race-id",
                "status": "sent",
            })
        return result

    buffer.update_one = write_receipt_then_save_response
    assert run(jobs.record_receipt(
        "branch-a", "meta_cloud", "race-id", "delivered",
        timestamp="2026-01-01T01:00:00+00:00", tenant_slug="tenant-a",
    )) == 1
    assert db["whatsapp_campaign_job_items"].rows[0]["delivery_status"] == "delivered"
    assert not buffer.rows


def test_receipt_buffer_is_tenant_branch_provider_scoped_and_waha_alias_safe(monkeypatch):
    db = ReceiptDB()
    monkeypatch.setattr(jobs, "_db", db)
    run(jobs.record_receipt(
        "branch-a", "waha", "stanza-id", "read",
        timestamp="2026-01-01T02:00:00+00:00", tenant_slug="tenant-a",
    ))
    db["whatsapp_campaign_job_items"].rows.append({
        "id": "item-1", "branch_id": "branch-a", "provider": "waha",
        "provider_message_id": "canonical@id", "status": "sent",
    })

    # Wrong scope cannot consume the buffered receipt.
    assert run(jobs.reconcile_receipt(
        "branch-b", "waha", "canonical@id",
        aliases=["stanza-id"], tenant_slug="tenant-a",
    )) == 0
    assert run(jobs.reconcile_receipt(
        "branch-a", "meta_cloud", "canonical@id",
        aliases=["stanza-id"], tenant_slug="tenant-a",
    )) == 0
    assert run(jobs.reconcile_receipt(
        "branch-a", "waha", "canonical@id",
        aliases=["stanza-id"], tenant_slug="tenant-b",
    )) == 0

    assert run(jobs.reconcile_receipt(
        "branch-a", "waha", "canonical@id",
        aliases=["stanza-id"], tenant_slug="tenant-a",
    )) == 1
    assert db["whatsapp_campaign_job_items"].rows[0]["delivery_status"] == "read"


def test_failure_before_ack_and_late_receipts_remain_monotonic(monkeypatch):
    db = ReceiptDB([{
        "id": "item-1", "branch_id": "branch-a", "provider": "meta_cloud",
        "provider_message_id": "meta-media-response-id", "status": "sent",
    }])
    monkeypatch.setattr(jobs, "_db", db)

    run(jobs.record_receipt(
        "branch-a", "meta_cloud", "meta-media-response-id", "failed",
        timestamp="2026-01-01T00:00:00+00:00", error="temporary",
        tenant_slug="tenant-a",
    ))
    run(jobs.record_receipt(
        "branch-a", "meta_cloud", "meta-media-response-id", "accepted",
        timestamp="2026-01-01T00:30:00+00:00", tenant_slug="tenant-a",
    ))
    run(jobs.record_receipt(
        "branch-a", "meta_cloud", "meta-media-response-id", "read",
        timestamp="2026-01-01T02:00:00+00:00", tenant_slug="tenant-a",
    ))
    run(jobs.record_receipt(
        "branch-a", "meta_cloud", "meta-media-response-id", "delivered",
        timestamp="2026-01-01T01:00:00+00:00", tenant_slug="tenant-a",
    ))
    run(jobs.record_receipt(
        "branch-a", "meta_cloud", "meta-media-response-id", "accepted",
        timestamp="2026-01-01T00:15:00+00:00", tenant_slug="tenant-a",
    ))
    item = db["whatsapp_campaign_job_items"].rows[0]
    assert item["delivery_status"] == "read"
    assert item["accepted_at"] == "2026-01-01T00:30:00+00:00"
    assert item["read_at"] == "2026-01-01T02:00:00+00:00"
    assert item["error"] is None


def test_campaign_success_without_provider_id_is_unknown_and_not_resent(monkeypatch):
    calls = []

    async def no_id_meta_response(*args, **kwargs):
        calls.append(args)
        return True, None, None

    async def fence():
        return None

    monkeypatch.setattr(
        whatsapp_mod, "_send_meta_cloud_message_result", no_id_meta_response
    )
    with pytest.raises(RuntimeError, match="missing_message_id"):
        run(whatsapp_mod._dispatch_bulk_job_item(
            {
                "id": "item-1", "branch_id": "branch-a",
                "provider": "meta_cloud", "phone": "966500000001",
                "message": "campaign", "media_index": 0,
            },
            {},
            fence,
        ))
    assert len(calls) == 1


@pytest.mark.parametrize("receipt", ["delivered", "read"])
def test_report_groups_attachment_items_and_requires_receipts_for_delivery(monkeypatch, receipt):
    db = DB(
        [{"id": "job-1", "branch_id": "branch-a", "created_at": datetime.now(timezone.utc)}],
        [
            {
                "id": "item-1", "job_id": "job-1", "branch_id": "branch-a",
                "provider": "meta_cloud", "recipient_index": 0,
                "recipient_id": "member-1", "recipient_name": "A",
                "phone": "966500000001", "status": "sent",
                "provider_message_id": "wamid-delivered",
                "delivery_status": receipt, f"{receipt}_at": "2026-01-01T01:00:00+00:00",
            },
            {
                "id": "item-2", "job_id": "job-1", "branch_id": "branch-a",
                "provider": "meta_cloud", "recipient_index": 0,
                "recipient_id": "member-1", "recipient_name": "A",
                "phone": "966500000001", "status": "sent",
                # This attachment has provider acceptance only.
                "provider_message_id": "wamid-exact",
            },
        ],
    )
    monkeypatch.setattr(jobs, "_db", db)
    report = run(jobs.get_report("job-1", "branch-a"))

    assert report["summary"]["total"] == 1
    assert report["summary"]["partial"] == 1
    assert report["summary"]["delivered"] == 0
    assert report["summary"]["read"] == 0
    assert report["recipients"][0]["status"] == "partial"
    assert report["recipients"][0]["id"] == "member-1"
    assert report["recipients"][0]["delivery_status"] == receipt
    assert report["recipients"][0]["phone"] == "966500000001"


def test_exact_provider_receipt_updates_only_matching_campaign_item(monkeypatch):
    db = DB(
        [{"id": "job-3", "branch_id": "branch-a"}],
        [{
            "id": "item-1", "job_id": "job-3", "branch_id": "branch-a",
            "provider": "whatsflow", "recipient_index": 0,
            "phone": "966500000001", "status": "sent",
            "provider_message_id": "exact-provider-id",
        }],
    )
    monkeypatch.setattr(jobs, "_db", db)
    run(jobs.record_receipt(
        "branch-a", "whatsflow", "wrong-id", "read",
        timestamp="2026-01-01T01:00:00+00:00",
    ))
    assert "delivery_status" not in db["whatsapp_campaign_job_items"].rows[0]
    run(jobs.record_receipt(
        "branch-a", "whatsflow", "exact-provider-id", "read",
        timestamp="2026-01-01T01:00:00+00:00",
    ))
    report = run(jobs.get_report("job-3", "branch-a"))
    assert report["summary"]["read"] == 1
    assert report["recipients"][0]["read_at"] == "2026-01-01T01:00:00+00:00"


def test_provider_receipts_are_monotonic_and_idempotent(monkeypatch):
    db = DB(
        [{"id": "job-4", "branch_id": "branch-a"}],
        [{
            "id": "item-1", "job_id": "job-4", "branch_id": "branch-a",
            "provider": "waha", "recipient_index": 0,
            "phone": "966500000001", "status": "sent",
            "provider_message_id": "waha-id",
        }],
    )
    monkeypatch.setattr(jobs, "_db", db)
    run(jobs.record_receipt(
        "branch-a", "waha", "waha-id", "delivered",
        timestamp="2026-01-01T01:00:00+00:00",
    ))
    run(jobs.record_receipt(
        "branch-a", "waha", "waha-id", "sent",
        timestamp="2026-01-01T00:00:00+00:00",
    ))
    item = db["whatsapp_campaign_job_items"].rows[0]
    assert item["delivery_status"] == "delivered"
    assert item["delivered_at"] == "2026-01-01T01:00:00+00:00"
    run(jobs.record_receipt(
        "branch-a", "waha", "waha-id", "read",
        timestamp="2026-01-01T02:00:00+00:00",
    ))
    run(jobs.record_receipt(
        "branch-a", "waha", "waha-id", "delivered",
        timestamp="2026-01-01T03:00:00+00:00",
    ))
    assert item["delivery_status"] == "read"
    assert item["read_at"] == "2026-01-01T02:00:00+00:00"


def test_report_masks_phone_and_sanitizes_error_in_excel(monkeypatch):
    db = DB(
        [{
            "id": "job-2", "branch_id": "branch-a",
            "campaign_title": "=unsafe title",
            "created_at": datetime.now(timezone.utc),
        }],
        [{
            "id": "item-1", "job_id": "job-2", "branch_id": "branch-a",
            "provider": "meta_cloud", "recipient_index": 0,
            "phone": "966500000001", "status": "failed",
            "error": "provider failed for 966500000001\n=HYPERLINK('x')",
        }],
    )
    monkeypatch.setattr(jobs, "_db", db)
    report = run(jobs.get_report("job-2", "branch-a", phone_visible=False))
    assert report["recipients"][0]["phone"] == "********"
    assert "966500000001" not in report["recipients"][0]["error"]

    content = whatsapp_mod._campaign_report_xlsx(report)
    workbook = load_workbook(__import__("io").BytesIO(content), data_only=False)
    values = [cell.value for row in workbook["Summary"].iter_rows() for cell in row]
    assert "'=unsafe title" in values
    error = workbook["Recipients"]["I2"].value
    assert workbook["Recipients"]["I2"].data_type != "f"
    assert workbook["Summary"]["B2"].data_type != "f"
    leading_formula = {
        **report,
        "recipients": [{**report["recipients"][0], "error": "=HYPERLINK('x')"}],
    }
    formula_workbook = load_workbook(
        __import__("io").BytesIO(whatsapp_mod._campaign_report_xlsx(leading_formula)),
        data_only=False,
    )
    formula_cell = formula_workbook["Recipients"]["I2"]
    assert formula_cell.value.startswith("'=")
    assert formula_cell.data_type != "f"


def test_phone_hidden_report_uses_generic_error_and_unambiguous_names(monkeypatch):
    db = DB(
        [{"id": "job-5", "branch_id": "branch-a"}],
        [{
            "id": "item-1", "job_id": "job-5", "branch_id": "branch-a",
            "provider": "meta_cloud", "recipient_index": 0,
            "phone": "966500000001", "status": "failed",
            "error": "raw provider error: +966 500 000 001",
        }],
    )
    db["members"].rows.extend([
        {"branch_id": "branch-a", "phone": "966500000001", "name": "Zed"},
        {"branch_id": "branch-a", "phone": "966500000001", "name": "Amy"},
    ])
    monkeypatch.setattr(jobs, "_db", db)
    monkeypatch.setattr(whatsapp_mod, "_db", db)
    report = run(whatsapp_mod._campaign_report(
        "job-5", "branch-a",
        {"permissions": [], "user_id": "user-1"},
    ))
    recipient = report["recipients"][0]
    assert recipient["phone"] == "********"
    assert recipient["error"] == "Provider error"
    assert recipient["name"] == "Amy / Zed"


def test_empty_historical_report_is_explicitly_unavailable(monkeypatch):
    db = DB([{"id": "old", "branch_id": "branch-a"}], [])
    monkeypatch.setattr(jobs, "_db", db)
    report = run(jobs.get_report("old", "branch-a"))
    assert report["summary"]["total"] == 0
    assert report["recipients"] == []
    assert any("timestamps" in note for note in report["notes"])


def test_report_endpoint_enforces_bulk_access_and_branch_ownership(monkeypatch):
    denied = {"is_admin": False, "permissions": [], "branch_id": "branch-a"}
    with pytest.raises(Exception) as exc:
        run(whatsapp_mod.get_branch_cloud_job_report(
            "job", "branch-a", current_user=denied
        ))
    assert getattr(exc.value, "status_code", None) == 403

    wrong_branch = {
        "is_admin": False, "permissions": ["messages"], "branch_id": "branch-a",
    }
    with pytest.raises(Exception) as exc:
        run(whatsapp_mod.get_branch_cloud_job_report(
            "job", "branch-b", current_user=wrong_branch
        ))
    assert getattr(exc.value, "status_code", None) == 403
